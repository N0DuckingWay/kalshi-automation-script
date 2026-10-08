"""Tests for live_portfolio.py: reading the account from Kalshi, reading the
bot's trade log, matching the bot's purchases to Kalshi orders, the ledger of
every contract bought, sold and paid out, and what is worked out from it: the
daily prices, the account's value over time, each period's statistics, each
purchase's return, the holdings now, the check of each run's logged cash, the
whole view and its JSON log line.

All offline. historical._historical_get is replaced by FakeKalshi, which
serves each listing page by page by cursor, and by FakeCandles for the daily
candles, so no request ever leaves the machine. The trade logs are written by
reporter.append_to_prod_log itself, with its paths pointed at tmp_path (as
tests/test_reporter.py does) and its clock and the host's time zone pinned.
The ledger's cash is checked against a separate running-position formula over
300 random fill sequences, and the pair-weighted median against numpy's
median over 200 random sets of purchases.
"""
import copy
import dataclasses
import json
import logging
import math
import random
import re
import shutil
from collections import defaultdict
from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from itertools import pairwise
from types import SimpleNamespace

import numpy as np
import openpyxl
import pandas as pd
import pytest
from kalshi_python_sync.exceptions import ApiException

from kalshi_betting import auth, config, dashboard, historical, live_portfolio, reporter, treasury
from kalshi_betting.live_portfolio import (
    OTHER_BETS,
    Account,
    BotLeg,
    BotTrade,
    CashFlow,
    Fill,
    History,
    Holding,
    LiveView,
    Market,
    Marks,
    Payout,
    RunStart,
    TradeReturn,
)
from kalshi_betting.reporter import TradeResult
from kalshi_betting.scanner import ApiMarket, CandidatePair
from kalshi_betting.strategy import TradeSpec

from .test_scheduler import _host_zone

D = Decimal
_API = "/trade-api/v2"
_KEYS = {
    "/historical/fills": "fills",
    "/portfolio/fills": "fills",
    "/portfolio/settlements": "settlements",
    "/portfolio/deposits": "deposits",
    "/portfolio/withdrawals": "withdrawals",
    "/portfolio/positions": "market_positions",
    "/markets": "markets",
    "/historical/markets": "markets",
}
# A fixed moment for tests that never compare with the clock
T0 = datetime(2026, 9, 28, 9, 0, tzinfo=UTC)


def _iso(moment: datetime) -> str:
    """A moment as Kalshi writes it: ISO text in UTC ending in Z."""
    return moment.astimezone(UTC).isoformat().replace("+00:00", "Z")


def _ago(seconds: float) -> datetime:
    """The moment `seconds` before the real clock now (for reads that compare with it)."""
    return datetime.now(UTC) - timedelta(seconds=seconds)


def fill_row(fill_id, ticker, when, book_side, count, yes_price, *, fee="0.0000",
             order_id=None, **extra) -> dict:
    """A fill record as Kalshi's fills listings send it."""
    yes = D(str(yes_price))
    row = {"fill_id": fill_id, "order_id": order_id or f"order-{fill_id}", "ticker": ticker,
           "created_time": _iso(when), "book_side": book_side,
           "count_fp": f"{D(str(count)):.2f}", "yes_price_dollars": f"{yes:.4f}",
           "no_price_dollars": f"{1 - yes:.4f}", "fee_cost": fee}
    row.update(extra)
    return row


class FakeKalshi:
    """
    Stands in for historical._historical_get: serves each listing page by page.

    rows maps a path (after /trade-api/v2) to its records, newest first as
    Kalshi lists them; balance is the /portfolio/balance reply. Every call is
    recorded, and on_balance (if set) runs before each balance read, so a test
    can change the account while it is being read.
    """

    def __init__(self):
        self.rows = {path: [] for path in _KEYS}
        self.balance = {"balance_breakdown": [{"exchange_index": 0, "balance_dollars": "100.0000"}],
                        "portfolio_value": 0}
        self.calls = []
        self.on_balance = None

    def __call__(self, client, path, **params):
        assert path.startswith(_API), path
        short = path[len(_API):]
        self.calls.append((short, dict(params)))
        if short == "/portfolio/balance":
            if self.on_balance:
                self.on_balance()
            return copy.deepcopy(self.balance)
        rows = self.rows[short]
        if params.get("tickers"):
            wanted = set(params["tickers"].split(","))
            rows = [r for r in rows if r["ticker"] in wanted]
        start = int(params.get("cursor") or 0)
        limit = params["limit"]
        nxt = start + limit
        return {_KEYS[short]: copy.deepcopy(rows[start:nxt]),
                "cursor": str(nxt) if nxt < len(rows) else ""}

    def count(self, path: str) -> int:
        """How many requests went to `path`."""
        return sum(1 for short, _ in self.calls if short == path)


@pytest.fixture
def kalshi(monkeypatch):
    """A FakeKalshi in place of historical._historical_get, and sleeps that only record."""
    fake = FakeKalshi()
    fake.sleeps = []
    fake.on_sleep = None

    def sleep(seconds):
        fake.sleeps.append(seconds)
        if fake.on_sleep:
            fake.on_sleep()

    monkeypatch.setattr(historical, "_historical_get", fake)
    monkeypatch.setattr(live_portfolio, "time", SimpleNamespace(sleep=sleep))
    return fake


# ---- reading the account ---------------------------------------------------

class TestReadAccount:
    def test_live_and_archived_fills_are_deduped_and_sorted(self, kalshi):
        kalshi.rows["/historical/fills"] = [
            fill_row("f1", "M1", _ago(3000), "bid", 2, "0.30"),
            fill_row("f2", "M1", _ago(4000), "bid", 1, "0.30", fee="0.0100"),
        ]
        kalshi.rows["/portfolio/fills"] = [
            fill_row("f3", "M2", _ago(1000), "ask", 3, "0.60"),
            fill_row("f2", "M1", _ago(4000), "bid", 1, "0.30", fee="0.0200"),
        ]
        account = live_portfolio.read_account(object())
        assert [f.fill_id for f in account.fills] == ["f2", "f1", "f3"]
        # The live copy of a fill is kept over its archived copy
        assert account.fills[0].fee == D("0.0200")
        assert account.changing is False
        assert account.warnings == ()
        assert kalshi.sleeps == []

    def test_a_fill_moving_to_the_archive_mid_read_is_kept(self, kalshi, monkeypatch):
        # Kalshi's archive cutoff moves while the fills are read: the old fill
        # leaves the live listing for the archive between the two reads
        kalshi.rows["/portfolio/fills"] = [fill_row("new", "M1", _ago(3600), "bid", 1, "0.30"),
                                           fill_row("old", "M1", _ago(3 * 86400), "bid", 2, "0.30")]
        moved = []

        def get(client, path, **params):
            reply = kalshi(client, path, **params)
            if path.endswith("fills") and params.get("limit") != 1 and not moved:
                moved.append(path)
                kalshi.rows["/historical/fills"].append(kalshi.rows["/portfolio/fills"].pop())
            return reply

        monkeypatch.setattr(historical, "_historical_get", get)
        account = live_portfolio.read_account(object())
        assert [f.fill_id for f in account.fills] == ["old", "new"]
        assert account.changing is False

    def test_a_no_purchase_is_read_from_book_side_alone(self, kalshi):
        # Kalshi sends a NO purchase with action "sell": only book_side is reliable
        kalshi.rows["/portfolio/fills"] = [
            fill_row("f2", "M1", _ago(500), "bid", 1, "0.30", action="sell", side="no"),
            fill_row("f1", "M1", _ago(600), "ask", 4, "0.25", action="sell", side="no"),
        ]
        account = live_portfolio.read_account(object())
        no_buy, yes_buy = account.fills
        assert no_buy.buys_yes is False
        assert (no_buy.count, no_buy.yes_price, no_buy.no_price) == (D("4.00"), D("0.25"), D("0.75"))
        assert yes_buy.buys_yes is True

    @pytest.mark.parametrize("book_side", [None, "", "yes", "buy"])
    def test_a_fill_without_a_book_side_is_refused(self, book_side):
        row = fill_row("f1", "M1", T0, "bid", 1, "0.30")
        row["book_side"] = book_side
        with pytest.raises(ValueError, match="book_side"):
            live_portfolio._fill(row)

    def test_cash_comes_from_every_shard(self, kalshi):
        kalshi.balance = {"balance": 9999, "balance_dollars": "99.99", "balance_breakdown": [
            {"exchange_index": 0, "balance_dollars": "10.1234", "balance": "77.0000"},
            {"exchange_index": 1, "balance": "5.0001"},
        ]}
        account = live_portfolio.read_account(object())
        assert account.cash == D("15.1235")
        assert account.shards == 2
        assert account.warnings == ()

    def test_a_shard_listed_twice_counts_once(self, kalshi):
        # auth._balance_cents_by_shard keeps a repeated shard's last entry; so does this
        kalshi.balance = {"balance_breakdown": [
            {"exchange_index": 0, "balance_dollars": "5.0000"},
            {"exchange_index": 1, "balance_dollars": "2.0000"},
            {"exchange_index": 0, "balance_dollars": "6.0000"},
        ]}
        account = live_portfolio.read_account(object())
        assert (account.cash, account.shards) == (D("8.0000"), 2)
        assert account.warnings == (
            "Kalshi's balance reply lists shard 0 twice: only its last entry is counted",)
        assert auth._balance_cents_by_shard(kalshi.balance) == {0: 600, 1: 200}

    @pytest.mark.parametrize("balance, cash", [
        ({"balance_dollars": "7.5000", "balance": 1}, D("7.5000")),
        ({"balance": 1234}, D("12.34")),
        ({"balance": 1234, "balance_breakdown": []}, D("12.34")),
    ])
    def test_cash_without_shards_is_one_shard(self, balance, cash):
        assert live_portfolio._cash(balance) == (cash, 1, [])

    def test_unreadable_cash_is_refused(self):
        with pytest.raises(ValueError):
            live_portfolio._cash({"portfolio_value": 10})
        with pytest.raises(ValueError):
            live_portfolio._cash({"balance_breakdown": [{"exchange_index": 0}]})

    @pytest.mark.parametrize("reply, value", [
        ({"portfolio_value": 7850}, D("78.50")),
        ({"portfolio_value_dollars": "78.5099", "portfolio_value": 1}, D("78.50")),
        ({}, None),
        ({"portfolio_value": -5}, None),
    ])
    def test_the_positions_value_is_auths_reading(self, kalshi, reply, value):
        kalshi.balance = {"balance": 100, **reply}
        account = live_portfolio.read_account(object())
        assert account.kalshi_positions_value == value

    def test_the_positions_value_goes_through_auth_at_call_time(self, kalshi, monkeypatch):
        seen = []

        def reading(data):
            seen.append(data)
            return 4321

        monkeypatch.setattr(auth, "_positions_value_cents", reading)
        account = live_portfolio.read_account(object())
        assert account.kalshi_positions_value == D("43.21")
        assert seen == [kalshi.balance]

    def test_deposits_and_withdrawals(self, kalshi):
        kalshi.rows["/portfolio/deposits"] = [
            {"status": "applied", "amount_cents": 10000, "fee_cents": 25,
             "created_ts": int(T0.timestamp()), "finalized_ts": int(T0.timestamp()) + 60},
            {"status": "pending", "amount_cents": 5000, "fee_cents": 0,
             "created_ts": int(T0.timestamp())},
        ]
        kalshi.rows["/portfolio/withdrawals"] = [
            {"status": "applied", "amount_cents": 5000, "fee_cents": 100,
             "created_ts": int(T0.timestamp()) - 3600},
        ]
        account = live_portfolio.read_account(object())
        assert [(f.time, f.amount) for f in account.flows] == [
            (T0 - timedelta(hours=1), D("-51")),
            (T0 + timedelta(seconds=60), D("99.75")),
        ]

    def test_settlement_revenue_and_value_are_read_in_cents(self, kalshi):
        kalshi.rows["/portfolio/settlements"] = [
            {"ticker": "M1", "settled_time": _iso(T0), "market_result": "yes",
             "revenue": 1300, "value": 100, "fee_cost": "0.3100"},
            {"ticker": "M2", "settled_time": _iso(T0), "market_result": "scalar",
             "revenue": 0, "value": 35},
        ]
        account = live_portfolio.read_account(object())
        first, second = account.payouts
        assert (first.ticker, first.result, first.revenue, first.yes_value) == (
            "M1", "yes", D("13.00"), D("1.00"))
        assert (second.result, second.revenue, second.yes_value) == ("scalar", D("0"), D("0.35"))

    def test_positions_are_signed_and_zero_is_left_out(self, kalshi):
        kalshi.rows["/portfolio/positions"] = [
            {"ticker": "M1", "position_fp": "-13.00"},
            {"ticker": "M2", "position_fp": "4.00"},
            {"ticker": "M3", "position_fp": "0.00"},
        ]
        account = live_portfolio.read_account(object())
        assert account.positions == {"M1": D("-13.00"), "M2": D("4.00")}
        assert ("/portfolio/positions", {"cursor": None, "limit": config.LIVE_PAGE_SIZE,
                                         "count_filter": "position"}) in kalshi.calls

    def test_every_listing_is_read_page_by_page(self, kalshi, monkeypatch):
        monkeypatch.setattr(config, "LIVE_PAGE_SIZE", 2)
        kalshi.rows["/portfolio/fills"] = [
            fill_row(f"f{i}", "M1", _ago(1000 + i), "bid", 1, "0.30") for i in range(5)]
        account = live_portfolio.read_account(object())
        assert len(account.fills) == 5
        cursors = [p["cursor"] for path, p in kalshi.calls
                   if path == "/portfolio/fills" and p["limit"] == 2]
        assert cursors == [None, "2", "4"]

    def test_the_records_are_read_before_the_balance_and_positions(self, kalshi):
        live_portfolio.read_account(object())
        order = [path for path, params in kalshi.calls if params.get("limit") != 1]
        # Live fills before archived ones: a fill moving to the archive while
        # they are read is then in one of the two reads
        assert order == ["/portfolio/fills", "/historical/fills", "/portfolio/settlements",
                         "/portfolio/deposits", "/portfolio/withdrawals",
                         "/portfolio/balance", "/portfolio/positions"]


class TestPages:
    def test_a_repeated_cursor_raises(self, monkeypatch):
        cursors = iter(["A", "B", "A"])
        monkeypatch.setattr(historical, "_historical_get",
                            lambda client, path, **params: {"fills": [], "cursor": next(cursors)})
        with pytest.raises(ValueError, match="repeated a page cursor"):
            live_portfolio._pages(object(), "/portfolio/fills", "fills", limit=5)

    def test_a_listing_that_never_ends_raises(self, monkeypatch):
        numbers = iter(range(10_000))
        monkeypatch.setattr(config, "SCANNER_MAX_PAGES", 3)
        monkeypatch.setattr(historical, "_historical_get",
                            lambda client, path, **params: {"fills": [],
                                                            "cursor": str(next(numbers))})
        with pytest.raises(ValueError, match="did not finish within 3 pages"):
            live_portfolio._pages(object(), "/portfolio/fills", "fills", limit=5)

    @pytest.mark.parametrize("reply", [{"cursor": ""}, {"fills": None}, {"fills": "x"}])
    def test_a_reply_without_its_list_raises(self, monkeypatch, reply):
        monkeypatch.setattr(historical, "_historical_get", lambda client, path, **params: reply)
        with pytest.raises(ValueError, match="without a 'fills' list"):
            live_portfolio._pages(object(), "/portfolio/fills", "fills", limit=5)

    @pytest.mark.parametrize("reply", [[], "text", None])
    def test_a_reply_that_is_not_an_object_raises(self, monkeypatch, reply):
        monkeypatch.setattr(historical, "_historical_get", lambda client, path, **params: reply)
        with pytest.raises(ValueError, match="not an object"):
            live_portfolio._get(object(), "/portfolio/balance")

    def test_every_request_goes_to_the_signed_get_with_the_api_prefix(self, monkeypatch):
        calls = []

        def fake(client, path, **params):
            calls.append((client, path, params))
            return {"fills": [], "cursor": ""}

        monkeypatch.setattr(historical, "_historical_get", fake)
        client = object()
        live_portfolio._pages(client, "/portfolio/fills", "fills", limit=7)
        assert calls == [(client, "/trade-api/v2/portfolio/fills", {"cursor": None, "limit": 7})]


class TestReadsAgainWhenTheAccountChanges:
    def test_a_fill_arriving_mid_read_reads_again(self, kalshi):
        kalshi.rows["/portfolio/fills"] = [fill_row("f1", "M1", _ago(600), "bid", 1, "0.30")]

        def arrive():
            if kalshi.count("/portfolio/balance") == 1:
                kalshi.rows["/portfolio/fills"].insert(
                    0, fill_row("f2", "M1", _ago(300), "bid", 2, "0.31"))

        kalshi.on_balance = arrive
        account = live_portfolio.read_account(object())
        assert [f.fill_id for f in account.fills] == ["f1", "f2"]
        assert account.changing is False
        assert kalshi.count("/portfolio/balance") == 2
        assert kalshi.sleeps == [config.LIVE_READ_RETRY_PAUSE_SECONDS]

    def test_a_settlement_arriving_mid_read_reads_again(self, kalshi):
        def arrive():
            if kalshi.count("/portfolio/balance") == 1:
                kalshi.rows["/portfolio/settlements"].insert(0, {
                    "ticker": "M9", "settled_time": _iso(_ago(60)), "market_result": "no",
                    "revenue": 500})

        kalshi.on_balance = arrive
        account = live_portfolio.read_account(object())
        assert [p.ticker for p in account.payouts] == ["M9"]
        assert kalshi.count("/portfolio/balance") == 2
        assert account.changing is False

    def test_a_fill_under_the_settle_age_waits_and_reads_again(self, kalshi):
        kalshi.rows["/portfolio/fills"] = [fill_row("f1", "M1", _ago(1), "bid", 1, "0.30")]

        def age():                    # the pause lets the fill grow old enough
            kalshi.rows["/portfolio/fills"][0]["created_time"] = _iso(
                _ago(config.LIVE_READ_SETTLE_SECONDS + 5))

        kalshi.on_sleep = age
        account = live_portfolio.read_account(object())
        assert account.changing is False
        assert kalshi.count("/portfolio/balance") == 2
        assert kalshi.sleeps == [config.LIVE_READ_RETRY_PAUSE_SECONDS]

    def test_an_account_that_keeps_changing_is_marked(self, kalshi):
        def arrive():
            n = kalshi.count("/portfolio/balance")
            kalshi.rows["/portfolio/fills"].insert(
                0, fill_row(f"new{n}", "M1", _ago(600 - n), "bid", 1, "0.30"))

        kalshi.on_balance = arrive
        account = live_portfolio.read_account(object())
        assert account.changing is True
        assert kalshi.count("/portfolio/balance") == config.LIVE_READ_ATTEMPTS
        # A pause between reads, none after the last
        assert kalshi.sleeps == [config.LIVE_READ_RETRY_PAUSE_SECONDS] * (
            config.LIVE_READ_ATTEMPTS - 1)

    def test_fills_sharing_the_newest_moment_are_not_a_change(self, kalshi):
        # One order that walked three levels: three fills at one moment, listed
        # in an order that is not their ids' order
        when = _ago(900)
        kalshi.rows["/portfolio/fills"] = [
            fill_row("fill-a", "M1", when, "bid", 1, "0.30", order_id="o1"),
            fill_row("fill-c", "M1", when, "bid", 1, "0.31", order_id="o1"),
            fill_row("fill-b", "M1", when, "bid", 1, "0.32", order_id="o1"),
        ]
        account = live_portfolio.read_account(object())
        assert account.changing is False
        assert kalshi.count("/portfolio/balance") == 1
        assert kalshi.sleeps == []


class TestReadMarkets:
    def test_the_outcome_label_is_the_subtitle_else_yes_sub_title(self, kalshi):
        kalshi.rows["/markets"] = [
            {"ticker": "R1", "title": "When?", "yes_sub_title": "Before Dec 1, 2026"},
            {"ticker": "R2", "title": "When?", "subtitle": "Before Nov 1, 2026",
             "yes_sub_title": "ignored"},
            {"ticker": "R3", "title": "When?", "yes_sub_title": 7},
            {"ticker": "R4", "title": "When?"}]
        found = live_portfolio.read_markets(object(), ["R1", "R2", "R3", "R4"])
        assert {t: m.subtitle for t, m in found.items()} == {
            "R1": "Before Dec 1, 2026", "R2": "Before Nov 1, 2026", "R3": "", "R4": ""}

    def test_live_first_then_the_archive(self, kalshi, monkeypatch):
        monkeypatch.setattr(config, "LIVE_MARKET_TICKERS_PER_REQUEST", 2)
        kalshi.rows["/markets"] = [
            {"ticker": "M1", "event_ticker": "E1", "title": "One?", "status": "active",
             "yes_bid_dollars": "0.40", "yes_ask_dollars": "0.45", "last_price_dollars": "0.42",
             "result": ""},
            {"ticker": "M3", "event_ticker": "E3", "status": "active", "yes_bid_dollars": "x"},
        ]
        kalshi.rows["/historical/markets"] = [
            {"ticker": "M2", "event_ticker": "E2", "title": "Two?", "status": "finalized",
             "result": "scalar", "settlement_ts": _iso(T0), "settlement_value_dollars": "0.3500"},
            {"ticker": "M1", "event_ticker": "E1", "title": "Old copy", "status": "finalized"},
        ]
        found = live_portfolio.read_markets(object(), ["M2", "M1", "M4", "M3", "M1"])
        assert set(found) == {"M1", "M2", "M3"}
        assert found["M1"] == Market("M1", "E1", "One?", "active", "", None, None,
                                     D("0.40"), D("0.45"), D("0.42"))
        assert found["M2"].settled_at == T0
        # Only a market the live listing lacked is marked as the archive's
        assert found["M2"].archived and not found["M1"].archived and not found["M3"].archived
        assert (found["M2"].result, found["M2"].yes_value) == ("scalar", D("0.3500"))
        assert found["M3"].title == "M3" and found["M3"].yes_bid is None
        asked = [(path, params["tickers"]) for path, params in kalshi.calls]
        # Two per request; the archive is asked only for what the live listing lacked
        assert asked == [("/markets", "M1,M2"), ("/markets", "M3,M4"),
                         ("/historical/markets", "M2,M4")]

    def test_nothing_wanted_sends_nothing(self, kalshi):
        assert live_portfolio.read_markets(object(), []) == {}
        assert kalshi.calls == []


# ---- the bot's trade log ---------------------------------------------------

def _trade_result(ticker_a: str, ticker_b: str, status: str, *, pair_type: str = "time_series",
                  count: int = 13) -> TradeResult:
    """A TradeResult as trader.execute_trades hands reporter one."""
    def market(ticker):
        return ApiMarket(ticker=ticker, event_ticker=f"EV-{ticker}", title=f"Will {ticker}?",
                         subtitle="", status="active", close_time=None)
    pair = CandidatePair(market_a=market(ticker_a), market_b=market(ticker_b), pA=0.30, pB=0.60,
                         nA=0.70, tradeable=True, canonical_title="pair", pair_type=pair_type,
                         nB=0.40)
    spec = TradeSpec(pair=pair, x=count, y=count, total_cost=9.10, total_cost_with_fees=9.41,
                     min_payoff=3.0, profit_ratio=0.3, days_to_close=20,
                     monthly_profit_ratio=0.4, kelly_p=0.7, kelly_fraction=0.1)
    return TradeResult(spec=spec, status=status,
                       error=None if status == "executed" else "YES leg error: killed")


@pytest.fixture
def log_paths(tmp_path, monkeypatch):
    """Point reporter's trade log, its lock and its fallback folder at tmp_path."""
    log_path = tmp_path / "trade_log.xlsx"
    monkeypatch.setattr(reporter, "PROD_LOG_PATH", log_path)
    monkeypatch.setattr(reporter, "_LOCK_PATH", tmp_path / "trade_log.xlsx.lock")
    monkeypatch.setattr(reporter, "PROJECT_ROOT", tmp_path)
    return log_path


@pytest.fixture
def pacific():
    """The host's clock on Los Angeles time (as the bot's Mac keeps it), restored afterwards."""
    with _host_zone("America/Los_Angeles"):
        yield


def _write_run(monkeypatch, at: datetime, results: list, cash_before: float) -> None:
    """Have reporter append one run to the trade log as if its clock read `at`."""
    class _Clock(datetime):
        @classmethod
        def now(cls, tz=None):
            return at if tz is None else at.astimezone(tz)

    monkeypatch.setattr(reporter, "datetime", _Clock)
    reporter.append_to_prod_log(results, cash_before, cash_before)


class TestReadTradeLogs:
    def test_runs_trades_and_windows(self, log_paths, pacific, monkeypatch):
        run1, dry, run2 = T0, T0 + timedelta(hours=1), T0 + timedelta(hours=2)
        _write_run(monkeypatch, run1, [
            _trade_result("A1", "B1", "executed"),
            _trade_result("A2", "B2", "rolled_back"),
            _trade_result("A3", "B3", "executed", pair_type="same_title", count=4),
        ], 132.45)
        _write_run(monkeypatch, dry, [_trade_result("A1", "B1", "simulated")], 120.00)
        _write_run(monkeypatch, run2, [
            _trade_result("A4", "B4", "failed"),
            _trade_result("A5", "B5", "manual_review"),
        ], 110.07)
        trades, starts, warnings = live_portfolio.read_trade_logs(
            live_portfolio.trade_log_paths())
        assert warnings == []
        window = timedelta(seconds=config.LIVE_BOT_RUN_WINDOW_SECONDS)
        assert [(t.trade_id, t.status) for t in trades] == [
            ("trade_log.xlsx#3", "executed"), ("trade_log.xlsx#4", "rolled_back"),
            ("trade_log.xlsx#5", "executed"), ("trade_log.xlsx#10", "manual_review")]
        first, _, same_title, review = trades
        assert first.legs == (BotLeg("A1", "yes", D(13)), BotLeg("B1", "no", D(13)))
        assert same_title.legs == (BotLeg("A3", "no", D(4)), BotLeg("B3", "yes", D(4)))
        assert first.title == "Will A1?"
        assert (first.logged_at, first.run_after) == (run1, run1 - window)
        # The dry run in between does not move the second run's window
        assert (review.logged_at, review.run_after) == (run2, run1)
        assert [(s.logged_at, s.run_after, s.cash_before) for s in starts] == [
            (run1, run1 - window, D("132.45")), (run2, run1, D("110.07"))]

    def test_an_old_header_workbook_is_read(self, log_paths, pacific, monkeypatch):
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed")], 50.0)
        book = openpyxl.load_workbook(log_paths)
        sheet = book.active
        for column, header in ((12, "x — NO on A"), (13, "y — YES on B"),
                               (14, "Total Cost ($)"), (15, "Min Profit ($)")):
            sheet.cell(row=1, column=column, value=header)
        book.save(log_paths)
        trades, starts, warnings = live_portfolio.read_trade_logs([log_paths])
        assert warnings == []
        assert [t.legs for t in trades] == [(BotLeg("A1", "yes", D(13)), BotLeg("B1", "no", D(13)))]
        assert [s.cash_before for s in starts] == [D("50.00")]

    def test_a_row_repeated_in_a_fallback_copy_is_read_once(self, log_paths, pacific, monkeypatch):
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed")], 50.0)
        shutil.copy(log_paths, log_paths.parent / "trade_log_2026-09-28_020000_123456.xlsx")
        trades, starts, warnings = live_portfolio.read_trade_logs(
            live_portfolio.trade_log_paths())
        assert [t.trade_id for t in trades] == ["trade_log.xlsx#3"]
        assert len(starts) == 1 and warnings == []

    def test_a_fallback_copys_own_runs_are_read(self, log_paths, pacific, monkeypatch):
        run1, run2 = T0, T0 + timedelta(hours=1)
        _write_run(monkeypatch, run1, [_trade_result("A1", "B1", "executed")], 50.0)
        # The shared log's lock cannot be taken, so reporter writes run 2 to a fallback copy
        monkeypatch.setattr(reporter, "_acquire_lock", lambda path: None)
        _write_run(monkeypatch, run2, [_trade_result("A2", "B2", "rolled_back", count=7)], 40.0)
        paths = live_portfolio.trade_log_paths()
        assert [p.name for p in paths] == ["trade_log.xlsx",
                                           "trade_log_2026-09-28_030000_000000.xlsx"]
        trades, starts, warnings = live_portfolio.read_trade_logs(paths)
        assert warnings == []
        window = timedelta(seconds=config.LIVE_BOT_RUN_WINDOW_SECONDS)
        assert [(t.trade_id, t.status, t.legs, t.logged_at, t.run_after) for t in trades] == [
            ("trade_log.xlsx#3", "executed",
             (BotLeg("A1", "yes", D(13)), BotLeg("B1", "no", D(13))), run1, run1 - window),
            (f"{paths[1].name}#3", "rolled_back",
             (BotLeg("A2", "yes", D(7)), BotLeg("B2", "no", D(7))), run2, run1)]
        assert [(s.logged_at, s.run_after, s.cash_before) for s in starts] == [
            (run1, run1 - window, D("50.00")), (run2, run1, D("40.00"))]

    def test_date_and_time_cells_of_other_kinds_are_read(self, log_paths, pacific, monkeypatch):
        # A workbook saved by hand can turn the text cells into a date and a time
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed"),
                                     _trade_result("A2", "B2", "executed")], 50.0)
        book = openpyxl.load_workbook(log_paths)
        sheet = book.active
        sheet.cell(row=3, column=1, value=datetime(2026, 9, 28))
        sheet.cell(row=3, column=2, value=time(2, 0, 0))
        sheet.cell(row=4, column=1, value=date(2026, 9, 28))
        book.save(log_paths)
        trades, _, warnings = live_portfolio.read_trade_logs([log_paths])
        assert warnings == []
        assert [(t.trade_id, t.logged_at) for t in trades] == [
            ("trade_log.xlsx#3", T0), ("trade_log.xlsx#4", T0)]

    def test_a_status_with_spaces_is_read_and_an_unknown_one_warns(self, log_paths, pacific,
                                                                   monkeypatch):
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed"),
                                     _trade_result("A2", "B2", "executed")], 50.0)
        book = openpyxl.load_workbook(log_paths)
        sheet = book.active
        sheet.cell(row=3, column=17, value="executed ")
        sheet.cell(row=4, column=17, value="half-done")
        book.save(log_paths)
        trades, starts, warnings = live_portfolio.read_trade_logs([log_paths])
        assert [(t.trade_id, t.status) for t in trades] == [("trade_log.xlsx#3", "executed")]
        assert warnings == [f"trade_log.xlsx#4: its status 'half-done' is not one this page "
                            f"knows; its contracts count as {OTHER_BETS}"]
        assert len(starts) == 1

    def test_two_runs_in_the_repeated_hour_keep_their_order(self, log_paths, pacific,
                                                            monkeypatch):
        # On 2026-11-01 Los Angeles reads 01:00-01:59 twice. Run 1 logs at 01:50
        # the first time round, run 2 at 01:10 the second time, 20 minutes later
        run1 = datetime(2026, 11, 1, 8, 50, tzinfo=UTC)
        run2 = datetime(2026, 11, 1, 9, 10, tzinfo=UTC)
        _write_run(monkeypatch, run1, [_trade_result("A1", "B1", "executed")], 50.0)
        _write_run(monkeypatch, run2, [_trade_result("A2", "B2", "executed")], 40.0)
        trades, starts, warnings = live_portfolio.read_trade_logs([log_paths])
        assert warnings == []
        assert [(t.logged_at, t.run_after) for t in trades] == [
            (run1, run1 - timedelta(seconds=config.LIVE_BOT_RUN_WINDOW_SECONDS)), (run2, run1)]
        fills = [_fill("1b", "B1", run1 - timedelta(seconds=2), False, 13),
                 _fill("1a", "A1", run1 - timedelta(seconds=1), True, 13),
                 _fill("2b", "B2", run2 - timedelta(seconds=2), False, 13),
                 _fill("2a", "A2", run2 - timedelta(seconds=1), True, 13)]
        owner, unmatched = live_portfolio.match_bot_fills(trades, fills)
        assert owner == {"1a": "trade_log.xlsx#3", "1b": "trade_log.xlsx#3",
                         "2a": "trade_log.xlsx#5", "2b": "trade_log.xlsx#5"}
        assert unmatched == []

    def test_a_lone_run_in_the_repeated_hour_reads_as_the_later_time(self, log_paths, pacific,
                                                                     monkeypatch):
        # With nothing after it to say which 01:30 it was, the later one is
        # taken: its orders still come before it
        _write_run(monkeypatch, datetime(2026, 11, 1, 8, 30, tzinfo=UTC),
                   [_trade_result("A1", "B1", "executed")], 50.0)
        trades, _, _ = live_portfolio.read_trade_logs([log_paths])
        assert [t.logged_at for t in trades] == [datetime(2026, 11, 1, 9, 30, tzinfo=UTC)]

    def test_only_reporters_fallback_copies_are_read(self, tmp_path, log_paths):
        names = ["trade_log.xlsx", "trade_log copy.xlsx", "trade_log_mine.xlsx",
                 "trade_log_2026-09-28_020000_123456.xlsx",
                 "trade_log_2026-09-28_020000_123456-1.xlsx",
                 "trade_log_2026-09-28_020000_123456-0a1b2c3d.xlsx",
                 "trade_log_2026-09-28_020000.xlsx",
                 "trade_log_2026-09-28_020000_123456.xlsx.tmp"]
        for name in names:
            (tmp_path / name).write_bytes(b"")
        (tmp_path / "trade_log_2026-09-29_020000_123456.xlsx").mkdir()
        assert [p.name for p in live_portfolio.trade_log_paths()] == [
            "trade_log.xlsx", "trade_log_2026-09-28_020000_123456-0a1b2c3d.xlsx",
            "trade_log_2026-09-28_020000_123456-1.xlsx",
            "trade_log_2026-09-28_020000_123456.xlsx"]

    def test_no_trade_log_is_no_path(self, log_paths):
        assert live_portfolio.trade_log_paths() == []

    def test_local_time_is_read_as_utc(self, pacific):
        once = datetime(2026, 9, 28, 9, 27, tzinfo=UTC)
        assert live_portfolio._log_readings("2026-09-28", "02:27:00") == (once, once)
        assert live_portfolio._log_readings("2026-09-28", "02:27") == (once, once)
        winter = datetime(2026, 12, 1, 10, 27, tzinfo=UTC)
        assert live_portfolio._log_readings("2026-12-01", "02:27:00") == (winter, winter)

    def test_the_repeated_hour_has_two_readings(self, pacific):
        # 01:30 happens twice on 2026-11-01 in Los Angeles: 08:30 UTC, then 09:30 UTC
        assert live_portfolio._log_readings("2026-11-01", "01:30:00") == (
            datetime(2026, 11, 1, 8, 30, tzinfo=UTC), datetime(2026, 11, 1, 9, 30, tzinfo=UTC))
        before = datetime(2026, 11, 1, 7, 30, tzinfo=UTC)
        assert live_portfolio._log_readings("2026-11-01", "00:30:00") == (before, before)

    @pytest.mark.parametrize("day, clock", [("2026-09-28", None), ("", ""), ("bad", "02:27:00")])
    def test_an_unreadable_time_raises(self, day, clock):
        with pytest.raises(ValueError):
            live_portfolio._log_readings(day, clock)

    def test_a_damaged_file_is_a_warning(self, tmp_path, log_paths, pacific, monkeypatch):
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed")], 50.0)
        broken = tmp_path / "trade_log_2026-09-28_020000_123456.xlsx"
        broken.write_bytes(b"not a workbook")
        trades, _, warnings = live_portfolio.read_trade_logs(live_portfolio.trade_log_paths())
        assert [t.trade_id for t in trades] == ["trade_log.xlsx#3"]
        assert warnings == [f"Could not read {broken.name} (BadZipFile): its trades count as "
                            f"{OTHER_BETS}"]

    def test_a_workbook_that_is_not_a_trade_log_is_a_warning(self, tmp_path):
        path = tmp_path / "trade_log.xlsx"
        book = openpyxl.Workbook()
        book.active.append(["Something", "else"])
        book.save(path)
        assert live_portfolio.read_trade_logs([path]) == (
            [], [], [f"trade_log.xlsx is not a trade log this page can read: its trades count "
                     f"as {OTHER_BETS}"])

    def test_unreadable_rows_are_warnings(self, log_paths, pacific, monkeypatch):
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed"),
                                     _trade_result("A2", "B2", "executed")], 50.0)
        book = openpyxl.load_workbook(log_paths)
        sheet = book.active
        sheet.cell(row=3, column=2, value="25:99:00")       # time
        sheet.cell(row=4, column=18, value="no sides here")  # notes
        book.save(log_paths)
        trades, _, warnings = live_portfolio.read_trade_logs([log_paths])
        assert trades == []
        assert warnings == [
            "trade_log.xlsx row 3: unreadable date or time",
            f"trade_log.xlsx#4: its sides or counts cannot be read; its contracts count as "
            f"{OTHER_BETS}"]


# ---- which fills are the bot's ---------------------------------------------

def _fill(fill_id, ticker, when, buys_yes, count, yes_price="0.30", fee="0", order_id=None):
    """A Fill as read_account builds one."""
    yes = D(yes_price)
    return Fill(fill_id, f"order-{fill_id}" if order_id is None else order_id, ticker, when,
                buys_yes, D(count), yes, 1 - yes, D(fee))


def _bot(trade_id, legs, logged_at, *, status="executed", run_after=None):
    """A BotTrade: legs are (ticker, side, count) triples, market A's first."""
    return BotTrade(trade_id, status, trade_id, tuple(BotLeg(t, s, D(c)) for t, s, c in legs),
                    run_after or logged_at - timedelta(seconds=config.LIVE_BOT_RUN_WINDOW_SECONDS),
                    logged_at)


class TestMatchBotFills:
    def test_an_exact_order_is_taken_and_no_other(self):
        logged = T0 + timedelta(hours=1)
        trade = _bot("T1", [("A", "yes", 10), ("B", "no", 10)], logged)
        fills = [
            _fill("a1", "A", logged - timedelta(seconds=30), True, 6, order_id="oA"),
            _fill("a2", "A", logged - timedelta(seconds=30), True, 4, "0.31", order_id="oA"),
            _fill("b1", "B", logged - timedelta(seconds=31), False, 8, order_id="partial"),
            _fill("b2", "B", logged - timedelta(seconds=32), False, 12, order_id="oversized"),
            _fill("b3", "B", logged - timedelta(seconds=33), True, 10, order_id="other-way"),
        ]
        owner, warnings = live_portfolio.match_bot_fills([trade], fills)
        assert owner == {"a1": "T1", "a2": "T1"}
        assert warnings == [f"The bot's purchase of 10 NO on B matches no Kalshi order: those "
                            f"contracts count as {OTHER_BETS}"]

    def test_consecutive_runs_of_one_pair_each_keep_their_own_orders(self):
        legs = [("A", "yes", 5), ("B", "no", 5)]
        logged_n = T0
        logged_next = T0 + timedelta(hours=1)
        run_n = _bot("N", legs, logged_n)
        run_next = _bot("N+1", legs, logged_next, run_after=logged_n)
        fills = [
            _fill("nB", "B", logged_n - timedelta(seconds=31), False, 5),
            _fill("nA", "A", logged_n - timedelta(seconds=30), True, 5),
            _fill("xB", "B", logged_n + timedelta(seconds=60), False, 5),
            _fill("xA", "A", logged_n + timedelta(seconds=61), True, 5),
        ]
        owner, warnings = live_portfolio.match_bot_fills([run_next, run_n], fills)
        assert owner == {"nA": "N", "nB": "N", "xA": "N+1", "xB": "N+1"}
        assert warnings == []

    @pytest.mark.parametrize("after", [0.2, 0.4, 0.6, 0.999])
    def test_orders_inside_the_logged_second_stay_with_their_run(self, after):
        # The Time cell shows the whole second, rounded down, so a run's last
        # orders can land just after the time shown (0 to 0.7 s on the real
        # account). Two runs 30 minutes apart buy the same legs and sizes.
        legs = [("A", "yes", 5), ("B", "no", 5)]
        logged_n = datetime(2026, 10, 5, 16, 0, 5, tzinfo=UTC)
        logged_next = logged_n + timedelta(minutes=30)
        run_n = _bot("N", legs, logged_n)
        run_next = _bot("N+1", legs, logged_next, run_after=logged_n)
        fills = []
        for name, logged in (("n", logged_n), ("m", logged_next)):
            fills += [_fill(f"{name}B", "B", logged - timedelta(seconds=1.5), False, 5),
                      _fill(f"{name}A", "A", logged + timedelta(seconds=after), True, 5)]
        owner, warnings = live_portfolio.match_bot_fills([run_n, run_next], fills)
        assert owner == {"nA": "N", "nB": "N", "mA": "N+1", "mB": "N+1"}
        assert warnings == []

    def test_the_previous_runs_logged_second_is_not_the_next_runs(self):
        # The previous run (whose pair sent nothing that filled) logged at
        # 16:00:05; an order inside that second is not the next run's, whose
        # own order lands 30 s past its log time (its clock running behind)
        logged_prev = datetime(2026, 10, 5, 16, 0, 5, tzinfo=UTC)
        logged = logged_prev + timedelta(minutes=30)
        run = _bot("N+1", [("A", "yes", 5), ("B", "no", 5)], logged, run_after=logged_prev)
        fills = [_fill("yours", "A", logged_prev + timedelta(seconds=0.5), True, 5),
                 _fill("b", "B", logged - timedelta(seconds=2), False, 5),
                 _fill("a", "A", logged + timedelta(seconds=30), True, 5)]
        owner, warnings = live_portfolio.match_bot_fills([run], fills)
        assert owner == {"a": "N+1", "b": "N+1"}
        assert warnings == []

    def test_an_order_just_past_the_log_time_is_found(self):
        logged = T0
        trade = _bot("T1", [("A", "yes", 5), ("B", "no", 5)], logged)
        fills = [
            _fill("b", "B", logged + timedelta(seconds=90), False, 5),
            _fill("a", "A", logged + timedelta(seconds=91), True, 5),
            # Past the logged second and the clock allowance: never the bot's
            _fill("late", "A", logged + timedelta(
                seconds=config.LIVE_TRADE_LOG_TIME_STEP_SECONDS
                + config.LIVE_BOT_RUN_CLOCK_SLACK_SECONDS + 1), True, 5),
        ]
        owner, warnings = live_portfolio.match_bot_fills([trade], fills)
        assert owner == {"a": "T1", "b": "T1"}
        assert warnings == []

    def test_in_the_window_the_latest_order_wins_and_past_it_the_earliest(self):
        logged = T0
        inside = _bot("inside", [("A", "yes", 5)] * 2, logged)
        fills = [_fill("early", "A", logged - timedelta(minutes=30), True, 5),
                 _fill("late", "A", logged - timedelta(minutes=1), True, 5)]
        owner, _ = live_portfolio.match_bot_fills([inside], fills)
        assert owner == {"early": "inside", "late": "inside"}
        one_leg = BotTrade("one", "executed", "one",
                           (BotLeg("A", "yes", D(5)), BotLeg("C", "no", D(1))),
                           logged - timedelta(hours=1), logged)
        assert live_portfolio.match_bot_fills([one_leg], fills)[0] == {"late": "one"}
        past = [_fill("p1", "A", logged + timedelta(seconds=10), True, 5),
                _fill("p2", "A", logged + timedelta(seconds=20), True, 5)]
        assert live_portfolio.match_bot_fills([one_leg], past)[0] == {"p1": "one"}

    def test_a_rolled_back_row_claims_its_no_purchase_and_its_unwind(self):
        logged = T0
        rolled = _bot("R", [("A", "yes", 13), ("B", "no", 13)], logged, status="rolled_back")
        retry = _bot("S", [("A", "yes", 13), ("B", "no", 13)], logged + timedelta(seconds=90),
                     status="rolled_back", run_after=logged)
        fills = [
            _fill("r-no", "B", logged - timedelta(seconds=20), False, 13),
            _fill("someone", "A", logged - timedelta(seconds=19), True, 13),
            _fill("r-unwind", "B", logged - timedelta(seconds=18), True, 13, "0.70"),
            # The retry, within the first row's clock allowance: its own orders
            _fill("s-no", "B", logged + timedelta(seconds=60), False, 13),
            _fill("s-unwind", "B", logged + timedelta(seconds=62), True, 13, "0.70"),
        ]
        owner, warnings = live_portfolio.match_bot_fills([rolled, retry], fills)
        assert owner == {"r-no": "R", "r-unwind": "R", "s-no": "S", "s-unwind": "S"}
        assert warnings == []

    def test_an_executed_retry_keeps_its_own_orders(self):
        logged = T0
        rolled = _bot("R", [("A", "yes", 13), ("B", "no", 13)], logged, status="rolled_back")
        retry = _bot("S", [("A", "yes", 13), ("B", "no", 13)], logged + timedelta(seconds=90),
                     run_after=logged)
        fills = [
            _fill("r-no", "B", logged - timedelta(seconds=20), False, 13),
            _fill("r-unwind", "B", logged - timedelta(seconds=18), True, 13, "0.70"),
            _fill("s-no", "B", logged + timedelta(seconds=60), False, 13),
            _fill("s-yes", "A", logged + timedelta(seconds=61), True, 13),
        ]
        owner, warnings = live_portfolio.match_bot_fills([rolled, retry], fills)
        assert owner == {"r-no": "R", "r-unwind": "R", "s-no": "S", "s-yes": "S"}
        assert warnings == []

    def test_a_partial_unwind_claims_only_up_to_the_no_count(self):
        logged = T0
        failed = _bot("F", [("A", "yes", 10), ("B", "no", 10)], logged, status="rollback_failed")
        fills = [
            _fill("no", "B", logged - timedelta(seconds=20), False, 10),
            _fill("part", "B", logged - timedelta(seconds=19), True, 4, "0.70"),
            # Your own buy-back later the same minute: more than the bot's NO left
            _fill("yours", "B", logged + timedelta(seconds=30), True, 7, "0.70"),
        ]
        owner, warnings = live_portfolio.match_bot_fills([failed], fills)
        assert owner == {"no": "F", "part": "F"}
        assert warnings == []

    @pytest.mark.parametrize("first_unwound, retry_status, retry_unwound", [
        (0, "rolled_back", 10), (4, "rollback_failed", 6)])
    def test_a_retrys_unwind_stays_with_the_retry(self, first_unwound, retry_status,
                                                  retry_unwound):
        # The first row's unwind failed (wholly, or after 4 of 10); 90 s later a
        # retry rolls back. Its unwind lies inside the first row's allowance too.
        logged = T0
        legs = [("A", "yes", 10), ("B", "no", 10)]
        first = _bot("R1", legs, logged, status="rollback_failed")
        retry = _bot("R2", legs, logged + timedelta(seconds=90), status=retry_status,
                     run_after=logged)
        fills = [_fill("r1-no", "B", logged - timedelta(seconds=20), False, 10)]
        if first_unwound:
            fills.append(_fill("r1-unwind", "B", logged - timedelta(seconds=19), True,
                               first_unwound, "0.70"))
        fills += [_fill("r2-no", "B", logged + timedelta(seconds=60), False, 10),
                  _fill("r2-unwind", "B", logged + timedelta(seconds=62), True, retry_unwound,
                        "0.70")]
        owner, warnings = live_portfolio.match_bot_fills([first, retry], fills)
        expected = {"r1-no": "R1", "r2-no": "R2", "r2-unwind": "R2"}
        if first_unwound:
            expected["r1-unwind"] = "R1"
        assert owner == expected
        assert warnings == []

    def test_a_manual_review_row_claims_no_buy_back(self):
        # On manual review the bot sends no unwind: a bid on its NO market just
        # after the alert is your own (a flatten by hand), not the bot's
        review = _bot("MR", [("A", "yes", 7), ("B", "no", 7)], T0, status="manual_review")
        fills = [_fill("mr-no", "B", T0 - timedelta(seconds=20), False, 7),
                 _fill("flatten", "B", T0 + timedelta(seconds=60), True, 7, "0.70")]
        owner, warnings = live_portfolio.match_bot_fills([review], fills)
        assert owner == {"mr-no": "MR"}
        assert warnings == []

    def test_unmatched_rows_that_were_not_executed_are_silent(self):
        review = _bot("M", [("A", "yes", 3), ("B", "no", 3)], T0, status="manual_review")
        assert live_portfolio.match_bot_fills([review], []) == ({}, [])

    def test_a_fill_without_an_order_id_is_its_own_order(self):
        trade = _bot("T1", [("A", "yes", 5), ("B", "no", 3)], T0)
        fills = [_fill("x", "A", T0 - timedelta(seconds=9), True, 2, order_id=""),
                 _fill("y", "A", T0 - timedelta(seconds=9), True, 3, order_id=""),
                 _fill("z", "B", T0 - timedelta(seconds=9), False, 3, order_id="")]
        owner, _ = live_portfolio.match_bot_fills([trade], fills)
        assert owner == {"z": "T1"}

    def test_your_close_partners_are_your_sales(self):
        fills = [_fill("bot", "A", T0, True, 1),
                 _fill("bot-c", "C", T0, False, 2),
                 _fill("y1", "A", T0 + timedelta(minutes=1), False, 1),   # sells YES A
                 _fill("y2", "B", T0 + timedelta(minutes=5), True, 1),    # buys YES B
                 _fill("y3", "C", T0 + timedelta(minutes=8), True, 1),    # sells NO C
                 _fill("y4", "D", T0 + timedelta(minutes=30), True, 1)]   # buys YES D
        partners = live_portfolio.manual_partners(fills, {"bot": "T1", "bot-c": "T1"})
        # Only a sale (a fill that closes what the account holds) is a partner
        assert partners == {"y1": frozenset({"C"}), "y2": frozenset({"A", "C"}),
                            "y3": frozenset({"A"}), "y4": frozenset()}


class TestTradeLogToOwners:
    def test_rows_written_by_reporter_own_their_orders(self, log_paths, pacific, monkeypatch):
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed", count=13)], 50.0)
        trades, _, _ = live_portfolio.read_trade_logs(live_portfolio.trade_log_paths())
        fills = [_fill("no", "B1", T0 - timedelta(seconds=3), False, 13),
                 _fill("yes", "A1", T0 - timedelta(seconds=2), True, 13)]
        owner, warnings = live_portfolio.match_bot_fills(trades, fills)
        assert owner == {"no": "trade_log.xlsx#3", "yes": "trade_log.xlsx#3"}
        assert warnings == []


# ---- the ledger ------------------------------------------------------------

def _ledger(fills, owners=None, payouts=(), legs=None):
    """build_ledger over fills; owners maps fill_id -> trade_id (the rest are Other bets)."""
    return live_portfolio.build_ledger(fills, payouts, owners or {}, legs or {})


def _holdings(ledger) -> dict:
    """Contracts each (owner, ticker, side) holds after every event, without zeros."""
    held = defaultdict(D)
    for e in ledger.events:
        held[(e.owner, e.ticker, e.side)] += e.contracts
    return {key: count for key, count in held.items() if count}


def _tuples(ledger):
    """Each event as (owner, ticker, side, contracts, cash, spent, basis)."""
    return [(e.owner, e.ticker, e.side, e.contracts, e.cash, e.spent, e.basis)
            for e in ledger.events]


class TestLedger:
    def test_cash_for_each_kind_of_fill(self):
        t = [T0 + timedelta(minutes=i) for i in range(4)]
        ledger = _ledger([
            _fill("yes-open", "Y", t[0], True, 10, "0.30", "0.07"),
            _fill("yes-close", "Y", t[1], False, 10, "0.45", "0.05"),
            _fill("no-open", "N", t[2], False, 5, "0.60", "0.03"),
            _fill("no-close", "N", t[3], True, 5, "0.70", "0.02"),
        ])
        assert _tuples(ledger) == [
            (OTHER_BETS, "Y", "yes", D(10), D("-3.07"), D("3.07"), D("3.07")),
            (OTHER_BETS, "Y", "yes", D(-10), D("4.45"), D(0), D("-3.07")),
            (OTHER_BETS, "N", "no", D(5), D("-2.03"), D("2.03"), D("2.03")),
            (OTHER_BETS, "N", "no", D(-5), D("1.48"), D(0), D("-2.03")),
        ]
        assert ledger.held_now == {} and ledger.warnings == ()

    def test_a_fill_crossing_zero_is_split(self):
        ledger = _ledger([_fill("open", "M", T0, False, 5, "0.60"),
                          _fill("cross", "M", T0 + timedelta(minutes=1), True, 8, "0.70", "0.08")])
        assert _tuples(ledger)[1:] == [
            (OTHER_BETS, "M", "no", D(-5), D("1.45"), D(0), D("-2.00")),
            (OTHER_BETS, "M", "yes", D(3), D("-2.13"), D("2.13"), D("2.13")),
        ]
        assert ledger.held_now == {"M": D(3)}

    def test_fees_add_back_exactly_over_many_pieces(self):
        # Three lots closed and the rest opened: four pieces of one fill's fee
        opens = [_fill(f"o{i}", "M", T0 + timedelta(seconds=i), True, 1, "0.30") for i in range(3)]
        sell = _fill("sell", "M", T0 + timedelta(minutes=1), False, 6, "0.40", "0.07")
        ledger = _ledger([*opens, sell], {"o0": "T1", "o1": "T2"})
        pieces = [e for e in ledger.events if e.time == sell.time]
        # What each piece's contracts are worth at the fill's prices, less its cash
        fees = [(-e.contracts * sell.yes_price if e.side == "yes"
                 else -e.contracts * sell.no_price) - e.cash for e in pieces]
        assert fees == [D("0.011666666667")] * 3 + [D("0.034999999999")]
        assert sum(fees, D(0)) == D("0.07")
        assert [e.owner for e in pieces] == [OTHER_BETS, "T1", "T2", OTHER_BETS]
        assert ledger.held_now == {"M": D(-3)}

    def test_a_closed_lot_keeps_no_cost(self):
        ledger = _ledger([_fill("open", "M", T0, True, 3, "0.33", "0.01"),
                          *[_fill(f"c{i}", "M", T0 + timedelta(minutes=i + 1), False, 1, "0.50")
                            for i in range(3)]])
        assert sum((e.basis for e in ledger.events), D(0)) == 0
        closes = [e.basis for e in ledger.events[1:]]
        assert closes == [D("-0.333333333333"), D("-0.333333333334"), D("-0.333333333333")]
        assert ledger.held_now == {}

    def test_an_unwind_takes_its_own_lots_first(self):
        fills = [_fill("yours", "B", T0, False, 5),
                 _fill("bot-no", "B", T0 + timedelta(minutes=1), False, 5),
                 _fill("unwind", "B", T0 + timedelta(minutes=2), True, 5, "0.70")]
        ledger = _ledger(fills, {"bot-no": "T", "unwind": "T"})
        assert _holdings(ledger) == {(OTHER_BETS, "B", "no"): D(5)}

    def test_a_pair_sold_by_hand_comes_out_of_its_own_purchase(self):
        # Two bot pairs share market A; you sell the newer pair (A and B2) by hand
        fills = [_fill("t1-b", "B1", T0, False, 10),
                 _fill("t1-a", "A", T0 + timedelta(seconds=1), True, 10),
                 _fill("t2-b", "B2", T0 + timedelta(hours=1), False, 10),
                 _fill("t2-a", "A", T0 + timedelta(hours=1, seconds=1), True, 10),
                 _fill("sell-a", "A", T0 + timedelta(days=1), False, 10, "0.50"),
                 _fill("buy-b2", "B2", T0 + timedelta(days=1, minutes=4), True, 10, "0.80")]
        owners = {"t1-b": "T1", "t1-a": "T1", "t2-b": "T2", "t2-a": "T2"}
        legs = {"T1": frozenset({"A", "B1"}), "T2": frozenset({"A", "B2"})}
        ledger = _ledger(fills, owners, legs=legs)
        assert _holdings(ledger) == {("T1", "B1", "no"): D(10), ("T1", "A", "yes"): D(10)}

    def test_without_a_partner_sale_the_oldest_purchase_goes_first(self):
        fills = [_fill("t1-a", "A", T0, True, 10),
                 _fill("t2-a", "A", T0 + timedelta(hours=1), True, 10),
                 _fill("sell-a", "A", T0 + timedelta(days=1), False, 10, "0.50")]
        legs = {"T1": frozenset({"A", "B1"}), "T2": frozenset({"A", "B2"})}
        ledger = _ledger(fills, {"t1-a": "T1", "t2-a": "T2"}, legs=legs)
        assert _holdings(ledger) == {("T2", "A", "yes"): D(10)}

    def test_a_purchase_beside_your_sale_does_not_pair_it(self):
        # You sell A, then 3 minutes later BUY NO on T2's other market: that is
        # no sale of T2's pair, so the sale closes the oldest purchase
        fills = [_fill("t1-a", "A", T0, True, 10),
                 _fill("t2-a", "A", T0 + timedelta(hours=1), True, 10),
                 _fill("sell-a", "A", T0 + timedelta(days=1), False, 10, "0.50"),
                 _fill("open-b2", "B2", T0 + timedelta(days=1, minutes=3), False, 5, "0.40")]
        legs = {"T1": frozenset({"A", "B1"}), "T2": frozenset({"A", "B2"})}
        ledger = _ledger(fills, {"t1-a": "T1", "t2-a": "T2"}, legs=legs)
        assert _holdings(ledger) == {("T2", "A", "yes"): D(10), (OTHER_BETS, "B2", "no"): D(5)}

    def test_your_sale_beside_a_purchase_takes_your_own_lots(self):
        # You sell the 4 YES A you bought yourself, then add 1 NO to B1, which
        # the bot holds NO of: the sale still comes out of your own lot
        fills = [_fill("t1-b", "B1", T0, False, 10),
                 _fill("t1-a", "A", T0 + timedelta(seconds=1), True, 10),
                 _fill("yours", "A", T0 + timedelta(hours=1), True, 4),
                 _fill("sell", "A", T0 + timedelta(days=1), False, 4, "0.50"),
                 _fill("more-b1", "B1", T0 + timedelta(days=1, minutes=2), False, 1, "0.40")]
        ledger = _ledger(fills, {"t1-b": "T1", "t1-a": "T1"},
                         legs={"T1": frozenset({"A", "B1"})})
        assert _holdings(ledger) == {("T1", "B1", "no"): D(10), ("T1", "A", "yes"): D(10),
                                     (OTHER_BETS, "B1", "no"): D(1)}

    def test_your_sale_takes_your_own_lots_before_the_bots(self):
        fills = [_fill("bot", "A", T0, True, 10),
                 _fill("yours", "A", T0 + timedelta(hours=1), True, 5),
                 _fill("sell", "A", T0 + timedelta(days=1), False, 7, "0.50")]
        ledger = _ledger(fills, {"bot": "T"}, legs={"T": frozenset({"A", "B"})})
        assert _holdings(ledger) == {("T", "A", "yes"): D(8)}

    @pytest.mark.parametrize("result, yes_value, paid_yes, paid_no", [
        ("yes", None, D(10), D(0)),
        ("no", None, D(0), D(4)),
        ("scalar", D("0.35"), D("3.50"), D("2.60")),
    ])
    def test_a_payout_pays_the_winning_side(self, result, yes_value, paid_yes, paid_no):
        fills = [_fill("y", "Y", T0, True, 10), _fill("n", "N", T0, False, 4)]
        settled = T0 + timedelta(days=1)
        payouts = [Payout("Y", settled, result, yes_value, None),
                   Payout("N", settled, result, yes_value, None)]
        ledger = _ledger(fills, {"y": "T"}, payouts)
        paid = {(e.ticker, e.owner): (e.contracts, e.cash, e.basis)
                for e in ledger.events if e.time == settled}
        assert paid == {("Y", "T"): (D(-10), paid_yes, D("-3.00")),
                        ("N", OTHER_BETS): (D(-4), paid_no, D("-2.80"))}
        assert ledger.held_now == {} and ledger.warnings == ()

    def test_revenue_that_disagrees_warns(self):
        fills = [_fill("y", "Y", T0, True, 10)]
        settled = T0 + timedelta(days=1)
        assert _ledger(fills, payouts=[Payout("Y", settled, "yes", None, D("10.00"))]).warnings == ()
        assert _ledger(fills, payouts=[Payout("Y", settled, "yes", None, D("9.99"))]).warnings == ()
        assert _ledger(fills, payouts=[Payout("Y", settled, "yes", None, D("9.98"))]).warnings == (
            "Y paid $9.98 but the contracts held add up to $10",)

    def test_an_unreadable_result_warns_and_keeps_the_contracts(self):
        ledger = _ledger([_fill("y", "Y", T0, True, 10)],
                         payouts=[Payout("Y", T0 + timedelta(days=1), "void", None, D(0))])
        assert ledger.held_now == {"Y": D(10)}
        assert ledger.warnings == ("Y settled as 'void' with no payout this page can read",)

    def test_a_fill_and_a_payout_at_one_moment_take_the_fill_first(self):
        ledger = _ledger([_fill("y", "Y", T0, True, 10)],
                         payouts=[Payout("Y", T0, "yes", None, D(10))])
        assert ledger.held_now == {} and ledger.warnings == ()


def _account(fills=(), payouts=(), positions=None, read_at=None) -> Account:
    """An Account built directly, as read_account would return it."""
    return Account(read_at or T0 + timedelta(days=10), D(100), 1, None, positions or {},
                   tuple(fills), tuple(payouts), (), False)


class TestRebuiltPayouts:
    def test_payouts_kalshi_no_longer_lists_are_rebuilt(self):
        settled = T0 + timedelta(days=2)
        fills = [_fill(f"f{t}", t, T0, True, 3) for t in ("M1", "M2", "M3", "M4", "M5", "M6")]
        fills.append(_fill("f-M6-close", "M6", T0 + timedelta(hours=1), False, 3))
        listed = Payout("M4", settled, "no", None, D(0))
        account = _account(fills, [listed], positions={"M5": D(3)})
        assert live_portfolio.tickers_needing_payout(account) == {"M1", "M2", "M3"}
        markets = {
            "M1": Market("M1", "E1", "One", "finalized", "yes", settled, D(1), None, None, None),
            "M3": Market("M3", "E3", "Three", "closed", "", None, None, None, None, None),
        }
        payouts, warnings = live_portfolio.all_payouts(account, markets)
        assert payouts == [listed, Payout("M1", settled, "yes", D(1), None)]
        assert warnings == [
            "Market M2 could not be looked up: its contracts count as still held",
            "Market M3 has no settlement this page can read: its contracts count as still held"]

    @pytest.mark.parametrize("result, settled", [
        ("void", T0 + timedelta(days=2)), ("", T0 + timedelta(days=2)), ("yes", None)])
    def test_a_found_market_with_no_readable_settlement_warns(self, result, settled):
        # Kalshi holds none of these contracts and lists no settlement for them
        account = _account([_fill("f", "OLD", T0, True, 3)])
        market = Market("OLD", "E", "Old", "finalized", result, settled, None, None, None, None)
        assert live_portfolio.all_payouts(account, {"OLD": market}) == (
            [], ["Market OLD has no settlement this page can read: its contracts count as "
                 "still held"])

    def test_a_settlement_after_the_read_is_not_rebuilt(self):
        account = _account([_fill("f", "M1", T0, True, 3)], read_at=T0 + timedelta(days=1))
        later = Market("M1", "E1", "One", "finalized", "yes", T0 + timedelta(days=2), D(1),
                       None, None, None)
        assert live_portfolio.all_payouts(account, {"M1": later}) == ([], [])


# ---- the ledger's cash against a running-position formula ------------------

def _oracle(fills, payouts):
    """
    Cash and holdings by a simple running position, independent of the ledger's lots.

    A bid closes NO held (paid its NO price) then opens YES (its YES price);
    an ask closes YES held (paid its YES price) then opens NO (its NO price);
    the fee comes off; a payout pays the position for its side.
    """
    position = defaultdict(D)
    cash = D(0)
    for fill in sorted(fills, key=lambda f: (f.time, f.fill_id)):
        held = position[fill.ticker]
        if fill.buys_yes:
            close = min(fill.count, max(-held, D(0)))
            cash += close * fill.no_price - (fill.count - close) * fill.yes_price
            position[fill.ticker] = held + fill.count
        else:
            close = min(fill.count, max(held, D(0)))
            cash += close * fill.yes_price - (fill.count - close) * fill.no_price
            position[fill.ticker] = held - fill.count
        cash -= fill.fee
    for payout in payouts:
        held = position[payout.ticker]
        yes = {"yes": D(1), "no": D(0), "scalar": payout.yes_value}[payout.result]
        cash += max(held, D(0)) * yes + max(-held, D(0)) * (1 - yes)
        position[payout.ticker] = D(0)
    return cash, {t: n for t, n in position.items() if n}


def _random_case(seed: int):
    """Random fills (and sometimes payouts) over three markets, with random owners."""
    rng = random.Random(seed)
    tickers = ["M1", "M2", "M3"]
    owners = ["T1", "T2", "T3", OTHER_BETS]
    moment = T0
    fills = []
    for i in range(rng.randint(1, 40)):
        moment += timedelta(seconds=rng.choice([0, 0, 1, 30, 300, 3600]))
        count = rng.choice([D(1), D(2), D(3), D(5), D(8), D("0.5"), D("2.25"), D("13")])
        fee = D(rng.randint(0, 300)) / 10000
        fills.append(_fill(f"f{i:04d}", rng.choice(tickers), moment, rng.random() < 0.5,
                           count, f"{rng.randint(1, 99) / 100:.2f}", str(fee)))
    owner_by_fill = {}
    for fill in fills:
        owner = rng.choice(owners)
        if owner != OTHER_BETS:
            owner_by_fill[fill.fill_id] = owner
    legs = {o: frozenset(rng.sample(tickers, 2)) for o in owners[:3]}
    payouts = []
    if rng.random() < 0.5:
        for ticker in rng.sample(tickers, rng.randint(1, 3)):
            result = rng.choice(["yes", "no", "scalar"])
            value = D(rng.randint(0, 100)) / 100 if result == "scalar" else None
            payouts.append(Payout(ticker, moment + timedelta(days=1), result, value, None))
    return fills, owner_by_fill, legs, payouts


class TestLedgerMatchesARunningPosition:
    @pytest.mark.parametrize("seed", range(300))
    def test_cash_and_holdings(self, seed):
        fills, owner_by_fill, legs, payouts = _random_case(seed)
        ledger = live_portfolio.build_ledger(fills, payouts, owner_by_fill, legs)
        cash, held = _oracle(fills, payouts)
        assert sum((e.cash for e in ledger.events), D(0)) == cash
        assert ledger.held_now == held
        if not payouts:
            signed = defaultdict(D)
            for fill in fills:
                signed[fill.ticker] += fill.count if fill.buys_yes else -fill.count
            assert ledger.held_now == {t: n for t, n in signed.items() if n}
        # Contracts an owner no longer holds keep no cost
        contracts, basis = defaultdict(D), defaultdict(D)
        for e in ledger.events:
            contracts[(e.owner, e.ticker, e.side)] += e.contracts
            basis[(e.owner, e.ticker, e.side)] += e.basis
        for key, count in contracts.items():
            assert count >= 0
            if count == 0:
                assert basis[key] == 0
        assert ledger.warnings == ()


# ---- prices ----------------------------------------------------------------

def _utc(*parts) -> datetime:
    """A UTC moment from its parts (year, month, day, hour, ...)."""
    return datetime(*parts, tzinfo=UTC)


# Just before the bot's first fill (T0), and a read nine days later
START = T0 - timedelta(microseconds=1)
READ = _utc(2026, 10, 7, 12)
# The daily candles' closes around them, and the first one read_marks asks for:
# midnight New York time starting the day before START's day (Sep 27, 04:00 UTC)
CLOSES = live_portfolio.day_ends(T0 - timedelta(days=3), READ)
FROM_TS = int(_utc(2026, 9, 27, 4).timestamp())
TO_TS = int(READ.timestamp())


def _candle(end: datetime, bid=None, ask=None, last=None, *, key="close_dollars") -> dict:
    """A daily candle as Kalshi sends it: each side's close as dollar text under `key`."""
    def side(price):
        return {key: None if price is None else f"{D(str(price)):.4f}"}
    return {"end_period_ts": int(end.timestamp()), "yes_bid": side(bid), "yes_ask": side(ask),
            "price": side(last)}


def _market(ticker, *, status="active", result="", settled_at=None, yes_value=None,
            bid=None, ask=None, last=None, event=None, archived=False) -> Market:
    """A Market as read_markets builds one (archived: found only in Kalshi's archive)."""
    def dec(value):
        return None if value is None else D(str(value))
    return Market(ticker, event or f"EV-{ticker}", f"Will {ticker}?", status, result, settled_at,
                  yes_value, dec(bid), dec(ask), dec(last), archived)


class FakeCandles:
    """
    Stands in for historical._historical_get on Kalshi's two candle listings.

    batch maps a ticker to its candles on GET /markets/candlesticks; archive
    maps a ticker to its candles on GET /historical/markets/{ticker}/candlesticks
    (a ticker it lacks answers 404). fail maps a ticker to what its archive
    request raises, and fail_batch, when set, is raised by every batch
    request; batch_reply, when set, is what every batch request answers.
    Each answer holds the candles that close inside the request's
    [start_ts, end_ts]. Every call is recorded.
    """

    def __init__(self):
        self.batch, self.archive, self.fail = {}, {}, {}
        self.fail_batch = None
        self.batch_reply = None
        self.calls = []

    def __call__(self, client, path, **params):
        assert path.startswith(_API), path
        short = path[len(_API):]
        self.calls.append((short, dict(params)))
        assert params["period_interval"] == config.LIVE_CANDLE_PERIOD_MINUTES
        lo, hi = params["start_ts"], params["end_ts"]

        def inside(rows):
            return [copy.deepcopy(c) for c in rows if lo <= c["end_period_ts"] <= hi]

        if short == "/markets/candlesticks":
            if self.fail_batch is not None:
                raise self.fail_batch
            if self.batch_reply is not None:
                return copy.deepcopy(self.batch_reply)
            return {"markets": [{"market_ticker": t, "candlesticks": inside(self.batch[t])}
                                for t in params["market_tickers"].split(",") if t in self.batch]}
        ticker = re.fullmatch(r"/historical/markets/(.+)/candlesticks", short).group(1)
        if ticker in self.fail:
            raise self.fail[ticker]
        if ticker not in self.archive:
            raise ApiException(status=404, reason="Not Found")
        return {"candlesticks": inside(self.archive[ticker])}

    def batches(self) -> list[dict]:
        """The parameters of each batch request."""
        return [params for short, params in self.calls if short == "/markets/candlesticks"]

    def archived(self) -> list[str]:
        """The archive paths asked, in order."""
        return [short for short, _ in self.calls if short.startswith("/historical/")]


@pytest.fixture
def candles(monkeypatch):
    """A FakeCandles in place of historical._historical_get."""
    fake = FakeCandles()
    monkeypatch.setattr(historical, "_historical_get", fake)
    return fake


def _daily(mid, *, until=None):
    """The daily points read_marks gives a market quoted at `mid` on every close it asks for."""
    return tuple((c, D(mid)) for c in CLOSES
                 if FROM_TS <= c.timestamp() <= TO_TS and (until is None or c <= until))


class TestDayEnds:
    def test_closes_across_both_clock_changes(self):
        # New York leaves daylight time on 2026-11-01 and goes back to it on
        # 2027-03-14: its midnight is 04:00 UTC in summer, 05:00 UTC in winter
        assert live_portfolio.day_ends(_utc(2026, 10, 30, 12), _utc(2026, 11, 2, 12)) == [
            _utc(2026, 10, 31, 4), _utc(2026, 11, 1, 4), _utc(2026, 11, 2, 5)]
        assert live_portfolio.day_ends(_utc(2027, 3, 13, 12), _utc(2027, 3, 16, 12)) == [
            _utc(2027, 3, 14, 5), _utc(2027, 3, 15, 4), _utc(2027, 3, 16, 4)]

    def test_a_close_at_the_start_or_the_end_is_left_out(self):
        close = _utc(2026, 9, 29, 4)
        assert live_portfolio.day_ends(close, close + timedelta(days=2)) == [
            close + timedelta(days=1)]
        assert live_portfolio.day_ends(close, close) == []


class TestPrices:
    @pytest.mark.parametrize("bid, ask, last, mid", [
        ("0.40", "0.50", "0.30", "0.45"),
        ("0", "0.50", None, "0.25"),        # no bid at all: half the ask
        ("0", "0.03", "0.50", "0.015"),     # ... before any last trade
        ("0.40", "1", None, "0.70"),        # no ask at all: halfway from the bid to 1
        ("0.97", "1", "0.50", "0.985"),     # ... before any last trade
        ("0", "1", "0.30", "0.30"),         # both sides empty: the last trade
        ("0", "0", "0.30", "0.30"),
        ("0.40", "1.20", None, None),
        ("0", "-0.10", None, None),
        ("0.50", "0.50", None, None),       # a bid at the ask
        ("0.60", "0.50", None, None),
        ("-0.10", "0.50", None, None),
        (None, "0.50", None, None),
        ("0.40", None, None, None),
        ("0", "1", "0", None),              # a last trade of 0 or 1 is no price
        ("0", "1", "1", None),
        (None, None, "0.62", "0.62"),
    ])
    def test_the_midpoint_refuses_kalshis_empty_sides(self, bid, ask, last, mid):
        def dec(value):
            return None if value is None else D(value)
        assert live_portfolio._mid(dec(bid), dec(ask), dec(last)) == dec(mid)

    def test_candles_in_both_shapes(self):
        end = _utc(2026, 9, 29, 4)
        new = _candle(end, "0.40", "0.50")
        old = _candle(end, "0.40", "0.50", key="close")
        assert live_portfolio._candle_mid(new) == (end, D("0.45"))
        assert live_portfolio._candle_mid(old) == (end, D("0.45"))
        # close_dollars is preferred when both are there
        both = {**new, "yes_bid": {"close_dollars": "0.4000", "close": 40}}
        assert live_portfolio._candle_mid(both) == (end, D("0.45"))

    def test_a_candle_with_no_quote_takes_its_last_trade(self):
        end = _utc(2026, 9, 29, 4)
        assert live_portfolio._candle_mid(_candle(end, "0", "1", "0.33")) == (end, D("0.33"))
        assert live_portfolio._candle_mid(_candle(end, "0", "1")) is None

    def test_a_day_with_no_trades(self):
        # Kalshi's candle for a day with no trades: the trades side carries only
        # the last price before the day, which is still the last trade
        end = _utc(2026, 9, 29, 4)
        quiet = {**_candle(end, "0.97", "1"), "price": {"previous_dollars": "0.9800"}}
        assert live_portfolio._candle_mid(quiet) == (end, D("0.985"))
        mirror = {**_candle(end, "0", "0.03"), "price": {"previous_dollars": "0.0200"}}
        assert live_portfolio._candle_mid(mirror) == (end, D("0.015"))
        empty = {**_candle(end, "0", "1"), "price": {"previous_dollars": "0.9800"}}
        assert live_portfolio._candle_mid(empty) == (end, D("0.98"))
        older = {**_candle(end, "0", "1"), "price": {"previous": "0.4100"}}
        assert live_portfolio._candle_mid(older) == (end, D("0.41"))
        # The day's own last trade comes before the price before it
        traded = {**_candle(end, "0", "1"), "price": {"close_dollars": "0.6000",
                                                      "previous_dollars": "0.9800"}}
        assert live_portfolio._candle_mid(traded) == (end, D("0.60"))
        assert live_portfolio._candle_mid(
            {**_candle(end, "0", "1"), "price": {"previous_dollars": "soon"}}) is None

    @pytest.mark.parametrize("change", [
        {"end_period_ts": None}, {"end_period_ts": "soon"}, {"end_period_ts": 10 ** 20},
        {"yes_ask": "0.50"}, {"yes_ask": {"close_dollars": "nan"}}])
    def test_an_unreadable_candle_gives_no_price(self, change):
        candle = {**_candle(_utc(2026, 9, 29, 4), "0.40", "0.50"), **change}
        assert live_portfolio._candle_mid(candle) is None

    def test_a_decided_result_comes_before_the_quote(self):
        quote = {"bid": "0.10", "ask": "0.20", "last": "0.15"}
        assert live_portfolio.quote_mid(_market("M", result="yes", **quote)) == 1
        assert live_portfolio.quote_mid(_market("M", result="no", **quote)) == 0
        assert live_portfolio.quote_mid(
            _market("M", result="scalar", yes_value=D("0.35"), **quote)) == D("0.35")
        assert live_portfolio.quote_mid(_market("M", result="scalar", **quote)) == D("0.15")
        assert live_portfolio.quote_mid(_market("M", **quote)) == D("0.15")
        assert live_portfolio.quote_mid(_market("M", bid="0", ask="1", last="0.12")) == D("0.12")
        assert live_portfolio.quote_mid(_market("M")) is None

    def test_marks_carry_the_last_value_forward(self):
        first, second = CLOSES[3], CLOSES[5]
        marks = Marks({"M": ((first, D("0.4")), (second, D("0.5")))}, {"M": D("0.7")}, READ)
        assert marks.at("M", first - timedelta(seconds=1)) is None
        assert marks.at("M", first) == D("0.4")
        assert marks.at("M", second - timedelta(seconds=1)) == D("0.4")
        assert marks.at("M", second + timedelta(days=1)) == D("0.5")
        assert marks.at("M", READ) == marks.at("M", READ + timedelta(days=1)) == D("0.7")
        assert marks.at("X", READ) is None
        # With no value now, the last daily price stands at the read too
        assert Marks(marks.daily, {}, READ).at("M", READ) == D("0.5")


class TestReadMarks:
    def test_the_batch_asks_few_enough_markets_and_candles(self, candles, monkeypatch):
        # The candle limit binds below the market limit here
        monkeypatch.setattr(config, "LIVE_CANDLE_TICKERS_PER_REQUEST", 5)
        monkeypatch.setattr(config, "LIVE_CANDLE_MAX_PER_REQUEST", 40)
        tickers = [f"M{i}" for i in range(7)]
        for ticker in tickers:
            candles.batch[ticker] = [_candle(c, "0.40", "0.50") for c in CLOSES]
        marks, warnings = live_portfolio.read_marks(
            object(), set(tickers), START, {t: _market(t) for t in tickers}, READ)
        assert warnings == []
        asked = candles.batches()
        # 12 candles a market fit 40 // 12 = 3 markets a request
        assert [p["market_tickers"] for p in asked] == ["M0,M1,M2", "M3,M4,M5", "M6"]
        assert {(p["start_ts"], p["end_ts"]) for p in asked} == {(FROM_TS, TO_TS)}
        for params in asked:
            assert (len(params["market_tickers"].split(","))
                    * live_portfolio._candles_in(params["start_ts"], params["end_ts"])) <= 40
        assert marks.daily == {t: _daily("0.45") for t in tickers}
        assert candles.archived() == []

    def test_a_long_span_is_read_in_windows(self, candles, monkeypatch):
        monkeypatch.setattr(config, "LIVE_CANDLE_MAX_PER_REQUEST", 5)
        candles.batch["M"] = [_candle(c, "0.40", "0.50") for c in CLOSES]
        marks, warnings = live_portfolio.read_marks(object(), {"M"}, START, {"M": _market("M")},
                                                    READ)
        windows = [(p["start_ts"], p["end_ts"]) for p in candles.batches()]
        assert len(windows) > 1 and warnings == []
        assert windows[0][0] == FROM_TS and windows[-1][1] == TO_TS
        assert all(later[0] == earlier[1] for earlier, later in pairwise(windows))
        assert all(live_portfolio._candles_in(lo, hi) <= 5 for lo, hi in windows)
        # A candle closing where two windows meet comes back twice and counts once
        assert marks.daily == {"M": _daily("0.45")}

    def test_a_long_span_of_many_markets_asks_few_enough_candles(self, candles, monkeypatch):
        monkeypatch.setattr(config, "LIVE_CANDLE_MAX_PER_REQUEST", 12)
        tickers = ["M0", "M1", "M2"]
        for ticker in tickers:
            candles.batch[ticker] = [_candle(c, "0.40", "0.50") for c in CLOSES]
        marks, warnings = live_portfolio.read_marks(
            object(), set(tickers), START, {t: _market(t) for t in tickers}, READ)
        assert warnings == []
        asked = [(p["market_tickers"], p["start_ts"], p["end_ts"]) for p in candles.batches()]
        # Ten days hold 12 candles of one market: one market a request; the last
        # eight hours hold 2, so all three fit in one request
        split = FROM_TS + 10 * 86_400
        assert asked == [("M0", FROM_TS, split), ("M1", FROM_TS, split), ("M2", FROM_TS, split),
                         ("M0,M1,M2", split, TO_TS)]
        for tickers_asked, lo, hi in asked:
            assert len(tickers_asked.split(",")) * live_portfolio._candles_in(lo, hi) <= 12
        assert marks.daily == {t: _daily("0.45") for t in tickers}

    def test_settled_markets_the_batch_lacks_come_from_the_archive(self, candles, tmp_path):
        settled = CLOSES[6] + timedelta(hours=3)
        markets = {
            "OPEN": _market("OPEN", bid="0.40", ask="0.50"),
            "FINAL": _market("FINAL", status="finalized", result="yes", settled_at=settled,
                             yes_value=D(1), archived=True),
            "DECIDED": _market("DECIDED", status="determined", result="no", settled_at=settled),
        }
        candles.batch["OPEN"] = [_candle(c, "0.40", "0.50") for c in CLOSES]
        for ticker in ("FINAL", "DECIDED"):
            candles.archive[ticker] = [_candle(c, "0.20", "0.30") for c in CLOSES if c <= settled]
        wanted = {"OPEN", "FINAL", "DECIDED", "GONE"}
        marks, warnings = live_portfolio.read_marks(object(), wanted, START, markets, READ)
        assert warnings == []
        # FINAL is only in the archive, so it goes straight there; GONE was
        # found nowhere, so it has no candles to ask for
        assert [p["market_tickers"] for p in candles.batches()] == ["DECIDED,OPEN"]
        assert candles.archived() == ["/historical/markets/DECIDED/candlesticks",
                                      "/historical/markets/FINAL/candlesticks"]
        assert marks.daily == {"OPEN": _daily("0.45"), "FINAL": _daily("0.25", until=settled),
                               "DECIDED": _daily("0.25", until=settled)}
        assert marks.now == {"OPEN": D("0.45"), "FINAL": D(1), "DECIDED": D(0)}
        # Only the finalized market is kept: never an open or not yet final one
        kept = config.LIVE_MARKS_CACHE_DIR
        assert kept.parent == tmp_path
        assert sorted(p.name for p in kept.iterdir()) == ["FINAL.json"]
        # The next read takes FINAL from its file, with no request for it
        candles.calls.clear()
        again, warnings = live_portfolio.read_marks(object(), wanted, START, markets, READ)
        assert again == marks and warnings == []
        assert [p["market_tickers"] for p in candles.batches()] == ["DECIDED,OPEN"]
        assert candles.archived() == ["/historical/markets/DECIDED/candlesticks"]

    def test_a_kept_file_that_starts_too_late_is_read_again(self, candles):
        settled = CLOSES[6] + timedelta(hours=3)
        markets = {"FINAL": _market("FINAL", status="finalized", result="yes",
                                    settled_at=settled, yes_value=D(1))}
        candles.archive["FINAL"] = [_candle(c, "0.20", "0.30") for c in CLOSES if c <= settled]
        live_portfolio.read_marks(object(), {"FINAL"}, START, markets, READ)
        candles.calls.clear()
        earlier = START - timedelta(days=2)
        marks, _ = live_portfolio.read_marks(object(), {"FINAL"}, earlier, markets, READ)
        assert candles.archived() == ["/historical/markets/FINAL/candlesticks"]
        kept = json.loads((config.LIVE_MARKS_CACHE_DIR / "FINAL.json").read_text())
        assert kept["start_ts"] == int(_utc(2026, 9, 25, 4).timestamp()) < FROM_TS
        # The first candle there is, Sep 26's, is before the first one read before
        assert marks.daily["FINAL"][0][0] == CLOSES[0] == _utc(2026, 9, 26, 4)

    def test_a_failed_archive_read_is_a_warning(self, candles):
        settled = CLOSES[6]
        markets = {"FINAL": _market("FINAL", status="finalized", result="yes",
                                    settled_at=settled, yes_value=D(1)),
                   "OLD": _market("OLD", status="finalized", result="no", settled_at=settled),
                   "OPEN": _market("OPEN")}
        candles.fail["FINAL"] = ApiException(status=503, reason="Service Unavailable")
        candles.archive["OLD"] = [_candle(c, "0.20", "0.30") for c in CLOSES if c <= settled]
        candles.batch["OPEN"] = [_candle(c, "0.40", "0.50") for c in CLOSES]
        marks, warnings = live_portfolio.read_marks(object(), set(markets), START, markets, READ)
        assert warnings == ["Daily prices of FINAL could not be read (HTTP 503 Service "
                            "Unavailable): it is valued at its last known price, or at what "
                            "it cost"]
        assert set(marks.daily) == {"OLD", "OPEN"}
        # FINAL is valued at its result now, and at what it cost before
        assert marks.at("FINAL", READ) == 1
        assert marks.at("FINAL", CLOSES[5]) is None
        assert not (config.LIVE_MARKS_CACHE_DIR / "FINAL.json").exists()

    def test_an_archive_answer_without_candles_is_a_warning(self, monkeypatch):
        settled = CLOSES[6]
        markets = {"FINAL": _market("FINAL", status="finalized", result="yes",
                                    settled_at=settled, yes_value=D(1))}
        monkeypatch.setattr(historical, "_historical_get", lambda client, path, **p: (
            {"markets": []} if path.endswith("/markets/candlesticks") else {"oops": 1}))
        marks, warnings = live_portfolio.read_marks(object(), {"FINAL"}, START, markets, READ)
        assert warnings == ["Daily prices of FINAL could not be read (ValueError: the archive's "
                            "candles for FINAL came without a candle list): it is valued at its "
                            "last known price, or at what it cost"]
        assert marks.daily == {}

    def test_a_failed_batch_request_is_a_warning(self, candles):
        settled = CLOSES[6]
        markets = {"FINAL": _market("FINAL", status="finalized", result="yes",
                                    settled_at=settled, yes_value=D(1)),
                   "OPEN": _market("OPEN", bid="0.40", ask="0.50")}
        candles.fail_batch = ApiException(status=500, reason="Internal Server Error")
        marks, warnings = live_portfolio.read_marks(object(), set(markets), START, markets, READ)
        assert warnings == ["Daily prices of 2 market(s) could not be read (HTTP 500 Internal "
                            "Server Error): they are valued at their last known price, or at "
                            "what they cost"]
        assert candles.archived() == []
        assert marks.daily == {} and marks.now == {"FINAL": D(1), "OPEN": D("0.45")}

    def test_a_failed_batch_request_never_costs_the_archives_markets(self, candles):
        settled = CLOSES[6]
        markets = {"OLD": _market("OLD", status="finalized", result="no", settled_at=settled,
                                  archived=True),
                   "OPEN": _market("OPEN", bid="0.40", ask="0.50")}
        candles.archive["OLD"] = [_candle(c, "0.20", "0.30") for c in CLOSES if c <= settled]
        candles.fail_batch = ApiException(status=400, reason="Bad Request")
        marks, warnings = live_portfolio.read_marks(object(), {"OLD", "OPEN", "GONE"}, START,
                                                    markets, READ)
        # The batch asked only for OPEN, and its failure is OPEN's alone
        assert [p["market_tickers"] for p in candles.batches()] == ["OPEN"]
        assert warnings == ["Daily prices of 1 market(s) could not be read (HTTP 400 Bad "
                            "Request): they are valued at their last known price, or at what "
                            "they cost"]
        assert candles.archived() == ["/historical/markets/OLD/candlesticks"]
        assert marks.daily == {"OLD": _daily("0.25", until=settled)}

    def test_an_open_market_the_batch_lacks_is_not_asked_of_the_archive(self, candles):
        markets = {"OPEN": _market("OPEN", bid="0.40", ask="0.50"), "QUIET": _market("QUIET")}
        candles.batch["OPEN"] = [_candle(c, "0.40", "0.50") for c in CLOSES]
        marks, warnings = live_portfolio.read_marks(object(), set(markets), START, markets, READ)
        assert warnings == [] and candles.archived() == []
        assert marks.daily == {"OPEN": _daily("0.45")}

    @pytest.mark.parametrize("reply, problem", [
        ({"oops": []}, "ValueError: the batch candle reply came without a market list"),
        ({"markets": 5}, "ValueError: the batch candle reply came without a market list"),
        ({"markets": None}, "ValueError: the batch candle reply came without a market list"),
    ])
    def test_a_batch_reply_without_a_market_list_is_a_warning(self, candles, reply, problem):
        markets = {"OPEN": _market("OPEN", bid="0.40", ask="0.50")}
        candles.batch_reply = reply
        marks, warnings = live_portfolio.read_marks(object(), {"OPEN"}, START, markets, READ)
        assert warnings == [f"Daily prices of 1 market(s) could not be read ({problem}): they "
                            f"are valued at their last known price, or at what they cost"]
        assert marks.daily == {}

    def test_entries_of_the_batch_reply_it_cannot_read_are_skipped(self, candles):
        markets = {"OPEN": _market("OPEN", bid="0.40", ask="0.50")}
        good = [_candle(c, "0.40", "0.50") for c in CLOSES if FROM_TS <= c.timestamp() <= TO_TS]
        candles.batch_reply = {"markets": [
            5, None, {"market_ticker": ["OPEN"], "candlesticks": good},
            {"market_ticker": {"t": "OPEN"}, "candlesticks": good},
            {"market_ticker": "OTHER", "candlesticks": good},
            {"market_ticker": "OPEN", "candlesticks": "none"},
            {"market_ticker": "OPEN", "candlesticks": [*good, 7]}]}
        marks, warnings = live_portfolio.read_marks(object(), {"OPEN"}, START, markets, READ)
        assert warnings == []
        assert marks.daily == {"OPEN": _daily("0.45")}

    def test_a_kept_file_for_another_market_is_not_used(self, candles):
        settled = CLOSES[6]
        markets = {"FINAL": _market("FINAL", status="finalized", result="yes",
                                    settled_at=settled, yes_value=D(1), archived=True)}
        candles.archive["FINAL"] = [_candle(c, "0.20", "0.30") for c in CLOSES if c <= settled]
        folder = config.LIVE_MARKS_CACHE_DIR
        folder.mkdir(parents=True)
        (folder / "FINAL.json").write_text(json.dumps({
            "ticker": "OTHER", "start_ts": FROM_TS, "end_ts": TO_TS,
            "candles": [_candle(c, "0.80", "0.90") for c in CLOSES]}))
        marks, warnings = live_portfolio.read_marks(object(), {"FINAL"}, START, markets, READ)
        assert warnings == []
        assert candles.archived() == ["/historical/markets/FINAL/candlesticks"]
        assert marks.daily == {"FINAL": _daily("0.25", until=settled)}
        # ... and the file is written again, for FINAL
        assert json.loads((folder / "FINAL.json").read_text())["ticker"] == "FINAL"

    def test_nothing_wanted_asks_nothing(self, candles):
        assert live_portfolio.read_marks(object(), set(), START, {}, READ) == (
            Marks({}, {}, READ), [])
        assert candles.calls == []

    def test_a_ticker_that_cannot_name_a_file_is_never_kept(self):
        assert live_portfolio._marks_file("../evil") is None
        assert live_portfolio._marks_file("a/b") is None
        assert live_portfolio._marks_file("KXBTCD-26SEP1517-T80999.99") == (
            config.LIVE_MARKS_CACHE_DIR / "KXBTCD-26SEP1517-T80999.99.json")


# ---- the account over time -------------------------------------------------

def _after(cash_at_start, ledger, *, flows=(), read_at=READ, shards=1) -> Account:
    """The account whose cash before every ledger event and flow was `cash_at_start`."""
    cash = (cash_at_start + sum((e.cash for e in ledger.events), D(0))
            + sum((f.amount for f in flows), D(0)))
    return Account(read_at, cash, shards, None, dict(ledger.held_now), (), (), tuple(flows), False)


def _crypto(owner: str) -> str:
    """Bot purchases T1, T2 ... are Crypto; everything else is Other bets."""
    return "Crypto" if owner.startswith("T") else OTHER_BETS


class TestHistory:
    @pytest.mark.parametrize("seed", range(50))
    def test_cash_rebuilt_backward_equals_cash_replayed_forward(self, seed):
        fills, owner_by_fill, legs, payouts = _random_case(seed)
        rng = random.Random(seed + 1000)
        flows = sorted((CashFlow(T0 + timedelta(hours=rng.randint(0, 60)),
                                 D(rng.randint(-5000, 5000)) / 100)
                        for _ in range(rng.randint(0, 4))), key=lambda f: f.time)
        ledger = live_portfolio.build_ledger(fills, payouts, owner_by_fill, legs)
        read_at = T0 + timedelta(days=3)
        start_cash = D("250.0000")
        account = _after(start_cash, ledger, flows=flows, read_at=read_at)
        changes = [(e.time, e.cash) for e in ledger.events] + [(f.time, f.amount) for f in flows]
        moments = sorted({when for when, _ in changes} | {START, read_at})
        for moment in moments:
            forward = start_cash + sum((amount for when, amount in changes if when <= moment),
                                       D(0))
            assert live_portfolio.cash_at(account, ledger, moment) == forward
        history = live_portfolio.build_history(account, ledger, Marks({}, {}, read_at), _crypto,
                                               START)
        assert history.cash == tuple(float(live_portfolio.cash_at(account, ledger, moment))
                                     for moment in history.times)
        assert sum(history.flows) == pytest.approx(float(sum((f.amount for f in flows), D(0))))

    def _half_sold_pair(self, sale_price: str) -> History:
        """A bot pair bought, then 5 of its 10 YES A sold by you the next day."""
        fills = [_fill("bot-b", "B", T0, False, 10, "0.60"),                  # NO B at 0.40
                 _fill("bot-a", "A", T0 + timedelta(seconds=1), True, 10, "0.30"),
                 _fill("sale", "A", _utc(2026, 9, 29, 10), False, 5, sale_price)]
        ledger = _ledger(fills, {"bot-b": "T1", "bot-a": "T1"},
                         legs={"T1": frozenset({"A", "B"})})
        read_at = _utc(2026, 10, 1, 12)
        closes = live_portfolio.day_ends(START, read_at)
        marks = Marks({"A": tuple((c, D("0.40")) for c in closes),
                       "B": tuple((c, D("0.60")) for c in closes)},
                      {"A": D("0.40"), "B": D("0.60")}, read_at)
        account = _after(D(100), ledger, read_at=read_at)
        return live_portfolio.build_history(account, ledger, marks, _crypto, START)

    def test_a_pair_half_sold_at_the_midpoint_leaves_the_total_alone(self):
        history = self._half_sold_pair("0.40")
        # Start, Sep 29, Sep 30 and Oct 1 closes, the read; the sale is on Sep 29 at 10:00
        assert len(history.times) == 5 and history.times[0] == START
        assert [history.total(i) for i in range(5)] == [100.0, 101.0, 101.0, 101.0, 101.0]
        value, net = history.value["Crypto"], history.net_cash["Crypto"]
        assert value[1] - value[2] == net[2] - net[1] == 2.0
        assert set(history.value) == {"Crypto"}

    def test_a_sale_below_the_midpoint_costs_only_the_gap(self):
        history = self._half_sold_pair("0.35")
        # 5 contracts sold 0.05 below the midpoint they were valued at
        assert history.total(2) - history.total(1) == pytest.approx(-0.25)
        value, net = history.value["Crypto"], history.net_cash["Crypto"]
        assert value[1] - value[2] == 2.0
        assert net[2] - net[1] == pytest.approx(1.75)

    def test_a_settlement_moves_value_into_cash(self):
        ledger = _ledger([_fill("m", "M", T0, True, 10, "0.30")],
                         payouts=[Payout("M", _utc(2026, 9, 29, 10), "yes", None, D(10))])
        read_at = _utc(2026, 10, 1, 12)
        closes = live_portfolio.day_ends(START, read_at)
        marks = Marks({"M": ((closes[0], D("0.70")),)}, {}, read_at)
        history = live_portfolio.build_history(_after(D(100), ledger, read_at=read_at), ledger,
                                               marks, _crypto, START)
        assert history.cash == (100.0, 97.0, 107.0, 107.0, 107.0)
        assert history.value == {OTHER_BETS: (0.0, 7.0, 0.0, 0.0, 0.0)}
        # Paid $10 for contracts valued at $7 the night before
        assert history.total(2) - history.total(1) == pytest.approx(3.0)
        assert history.net_cash[OTHER_BETS] == (0.0, -3.0, 7.0, 7.0, 7.0)
        assert history.spent[OTHER_BETS] == (0.0, 3.0, 3.0, 3.0, 3.0)
        assert history.steps[OTHER_BETS] == ((T0, -3.0), (_utc(2026, 9, 29, 10), 7.0))

    def test_a_group_with_no_value_but_cash_that_moved_is_kept(self):
        fills = [_fill("old-buy", "Z", T0 - timedelta(days=2), True, 1),
                 _fill("old-sell", "Z", T0 - timedelta(days=1), False, 1),
                 _fill("buy", "M", T0, True, 10, "0.30"),
                 _fill("sell", "M", T0 + timedelta(hours=1), False, 10, "0.35")]
        # T9's trade opened and closed before the start: nothing of it is shown
        ledger = _ledger(fills, {"old-buy": "T9", "old-sell": "T9"})
        read_at = _utc(2026, 10, 1, 12)
        history = live_portfolio.build_history(
            _after(D(100), ledger, read_at=read_at), ledger, Marks({}, {}, read_at),
            lambda owner: "Old" if owner == "T9" else OTHER_BETS, START)
        assert history.value == {OTHER_BETS: (0.0,) * 5}
        assert history.net_cash[OTHER_BETS][-1] == pytest.approx(0.5)

    def test_contracts_never_priced_are_valued_at_cost(self):
        ledger = _ledger([_fill("m", "M", T0, True, 10, "0.30", "0.07")])
        read_at = _utc(2026, 10, 1, 12)
        history = live_portfolio.build_history(_after(D(100), ledger, read_at=read_at), ledger,
                                               Marks({}, {}, read_at), _crypto, START)
        assert history.value[OTHER_BETS] == (0.0, 3.07, 3.07, 3.07, 3.07)
        assert [history.total(i) for i in range(5)] == pytest.approx([100.0] * 5)

    def test_a_deposit_is_valued_at_the_moment_it_lands(self):
        ledger = _ledger([_fill("m", "M", T0, True, 10, "0.30")])
        read_at = _utc(2026, 10, 1, 12)
        deposit = CashFlow(_utc(2026, 9, 29, 10), D("50"))
        account = _after(D(100), ledger, flows=[deposit], read_at=read_at)
        history = live_portfolio.build_history(account, ledger, Marks({}, {}, read_at), _crypto,
                                               START)
        assert history.times == (START, _utc(2026, 9, 29, 4), deposit.time,
                                 _utc(2026, 9, 30, 4), _utc(2026, 10, 1, 4), read_at)
        assert history.flows == (0.0, 0.0, 50.0, 0.0, 0.0, 0.0)
        assert history.cash == (100.0, 97.0, 147.0, 147.0, 147.0, 147.0)


# ---- statistics ------------------------------------------------------------

def _times(read_at):
    """The moments a history from START to `read_at` is valued at."""
    return (START, *live_portfolio.day_ends(START, read_at), read_at)


class TestPeriodStats:
    READ = _utc(2026, 10, 1, 12)

    def _deposit_history(self) -> History:
        """+10%, +10% with a $50 deposit, -5%, +2%: the last two whole days in the middle."""
        totals = (100.0, 110.0, 171.0, 162.45, 165.699)
        return History(_times(self.READ), totals, {}, {}, {}, {}, (0.0, 0.0, 50.0, 0.0, 0.0))

    def test_the_total_return_takes_a_deposit_out(self):
        stats = live_portfolio.period_stats(self._deposit_history(), [], None, "All", 0)
        assert stats.total_return == pytest.approx(1.1 * 1.1 * 0.95 * 1.02 - 1)
        # Profit leaves the deposit out
        assert stats.pnl == pytest.approx(165.699 - 100 - 50)
        assert (stats.label, stats.first, stats.last) == ("All", START, self.READ)

    def test_ratios_are_the_backtest_pages_over_whole_days_only(self):
        stats = live_portfolio.period_stats(self._deposit_history(), [], None, "All", 0)
        # The part-day before the first close and the one up to the read are left out
        whole = pd.Series([121.0 / 110.0 - 1, 162.45 / 171.0 - 1])
        assert stats.whole_days == 2
        assert stats.sharpe == dashboard._sharpe(whole, 0.0,
                                                 periods_per_year=config.CALENDAR_DAYS_PER_YEAR)
        assert stats.sortino == dashboard._sortino(whole, 0.0,
                                                   periods_per_year=config.CALENDAR_DAYS_PER_YEAR)
        every = pd.Series([110.0 / 100.0 - 1, *whole, 165.699 / 162.45 - 1])
        assert stats.sharpe != dashboard._sharpe(every, 0.0)

    def test_the_ratios_are_looked_up_when_called(self, monkeypatch):
        monkeypatch.setattr(dashboard, "_sharpe", lambda series, rf, *, periods_per_year: 7.0)
        monkeypatch.setattr(dashboard, "_sortino", lambda series, rf, *, periods_per_year: 8.0)
        stats = live_portfolio.period_stats(self._deposit_history(), [], None, "All", 0)
        assert (stats.sharpe, stats.sortino) == (7.0, 8.0)

    def test_the_hurdle_is_the_yield_on_the_share_in_positions(self):
        cash = (100.0, 50.0, 60.0, 70.0, 70.0)
        value = {"Crypto": (0.0, 60.0, 55.0, 40.0, 41.0)}            # totals 100 110 115 110 111
        zeros = {"Crypto": (0.0,) * 5}
        history = History(_times(self.READ), cash, value, zeros, zeros, {"Crypto": ()},
                          (0.0,) * 5)
        rates = treasury.RiskFreeRates(((date(2026, 9, 1), 0.04),), treasury.SOURCE_API, T0)
        stats = live_portfolio.period_stats(history, [], rates, "All", 0)
        whole = pd.Series([115.0 / 110.0 - 1, 110.0 / 115.0 - 1])
        rf = np.array([0.04, 0.04]) * np.array([1 - 50.0 / 110.0, 1 - 60.0 / 115.0])
        assert stats.sharpe == dashboard._sharpe(whole, rf, periods_per_year=365)
        assert stats.sortino == dashboard._sortino(whole, rf, periods_per_year=365)
        assert stats.sharpe != dashboard._sharpe(whole, 0.0, periods_per_year=365)

    def test_under_two_whole_days_there_are_no_ratios(self):
        read_at = _utc(2026, 9, 30, 12)                  # start, Sep 29, Sep 30, the read
        history = History(_times(read_at), (100.0, 101.0, 99.0, 100.0), {}, {}, {}, {},
                          (0.0,) * 4)
        stats = live_portfolio.period_stats(history, [], None, "All", 0)
        assert (stats.sharpe, stats.sortino, stats.whole_days) == (None, None, 1)
        assert stats.total_return == pytest.approx(0.0)

    def test_short_periods_on_a_short_history_are_all_of_it(self):
        read_at = T0 + timedelta(days=10)
        times = _times(read_at)
        rng = random.Random(7)
        cash = tuple(rng.uniform(40, 60) for _ in times)
        value = {"Crypto": tuple(rng.uniform(40, 60) for _ in times),
                 OTHER_BETS: tuple(rng.uniform(0, 5) for _ in times)}
        net = {g: tuple(rng.uniform(-5, 5) for _ in times) for g in value}
        steps = {g: tuple((t, rng.uniform(-9, 5)) for t in times[1:]) for g in value}
        history = History(times, cash, value, net, net, steps, (0.0,) * len(times))
        returns = [TradeReturn("T1", T0, D(4), 0.1, D(4), True),
                   TradeReturn("T2", T0 + timedelta(days=3), D(2), -0.2, D(0), False)]
        everything = live_portfolio.period_stats(history, returns, None, "All", 0)
        assert everything.purchases == 2 and everything.pairs == 6 and everything.open_pairs == 4
        for label, months in (("1M", 1), ("3M", 3), ("1Y", 12)):
            assert live_portfolio.period_stats(history, returns, None, label, months) == (
                dataclasses.replace(everything, label=label))

    def test_money_put_to_work_the_day_it_lands_earns_only_from_then(self):
        # $200: $100 in 200 YES A at 0.50. At 01:00 New York time on Sep 29,
        # $1,000 lands and buys 2,000 YES B at 0.50. Both are at 0.52 at the
        # next close: the day earned 3.67% on $1,200, not 22% on $200.
        landed = _utc(2026, 9, 29, 5)
        fills = [_fill("a", "A", T0, True, 200, "0.50"),
                 _fill("b", "B", landed + timedelta(minutes=1), True, 2000, "0.50")]
        ledger = _ledger(fills)
        read_at = _utc(2026, 9, 30, 12)
        closes = live_portfolio.day_ends(START, read_at)
        marks = Marks({"A": ((closes[0], D("0.50")), (closes[1], D("0.52"))),
                       "B": ((closes[1], D("0.52")),)}, {}, read_at)
        account = _after(D(200), ledger, flows=[CashFlow(landed, D(1000))], read_at=read_at)
        history = live_portfolio.build_history(account, ledger, marks, _crypto, START)
        assert history.times == (START, closes[0], landed, closes[1], read_at)
        assert [history.total(i) for i in range(5)] == pytest.approx(
            [200.0, 200.0, 1200.0, 1244.0, 1244.0])
        stats = live_portfolio.period_stats(history, [], None, "All", 0)
        assert stats.total_return == pytest.approx(1244 / 1200 - 1)
        assert stats.pnl == pytest.approx(44.0)
        # The whole day from the Sep 29 close to the Sep 30 close, chained over
        # the deposit, is the ratios' one day
        assert stats.whole_days == 1

    def test_a_whole_day_is_chained_over_the_moments_inside_it(self):
        times = (START, _utc(2026, 9, 29, 4), _utc(2026, 9, 29, 10), _utc(2026, 9, 30, 4),
                 _utc(2026, 9, 30, 15), _utc(2026, 10, 1, 4), _utc(2026, 10, 1, 12))
        # Sep 29: +5% up to a $50 deposit, then +10%; Sep 30: -5%, then -2%
        totals = (100.0, 110.0, 165.5, 182.05, 172.9475, 169.48855, 170.0)
        flows = (0.0, 0.0, 50.0, 0.0, 0.0, 0.0, 0.0)
        history = History(times, totals, {}, {}, {}, {}, flows)
        stats = live_portfolio.period_stats(history, [], None, "All", 0)
        steps = [110.0 / 100.0, 115.5 / 110.0, 182.05 / 165.5, 172.9475 / 182.05,
                 169.48855 / 172.9475, 170.0 / 169.48855]
        assert steps[1:5] == pytest.approx([1.05, 1.10, 0.95, 0.98])
        whole = pd.Series([steps[1] * steps[2] - 1, steps[3] * steps[4] - 1])
        assert stats.whole_days == 2
        assert stats.sharpe == pytest.approx(dashboard._sharpe(
            whole, 0.0, periods_per_year=config.CALENDAR_DAYS_PER_YEAR), rel=1e-12)
        assert stats.sortino == pytest.approx(dashboard._sortino(
            whole, 0.0, periods_per_year=config.CALENDAR_DAYS_PER_YEAR), rel=1e-12)
        assert stats.total_return == pytest.approx(math.prod(steps) - 1, rel=1e-12)

    def test_a_start_or_a_read_at_a_close_bounds_a_whole_day(self):
        # Both ends are daily closes: every day between them is whole
        start, read_at = _utc(2026, 9, 28, 4), _utc(2026, 10, 1, 4)
        times = (start, *live_portfolio.day_ends(start, read_at), read_at)
        history = History(times, (100.0, 101.0, 99.0, 102.0), {}, {}, {}, {}, (0.0,) * 4)
        stats = live_portfolio.period_stats(history, [], None, "All", 0)
        assert stats.whole_days == 3
        assert stats.sharpe == dashboard._sharpe(
            pd.Series([101.0 / 100.0 - 1, 99.0 / 101.0 - 1, 102.0 / 99.0 - 1]), 0.0,
            periods_per_year=config.CALENDAR_DAYS_PER_YEAR)

    def test_a_month_back_is_counted_on_new_york_s_calendar(self):
        # 22:00 New York time on Mar 30 is already Mar 31 in UTC: a month back is
        # Feb 28 at 22:00 New York time, so the period opens at the Mar 1 close
        start = _utc(2027, 1, 10, 12)
        read_at = _utc(2027, 3, 31, 2)
        times = (start, *live_portfolio.day_ends(start, read_at), read_at)
        history = History(times, (100.0,) * len(times), {}, {}, {}, {}, (0.0,) * len(times))
        assert live_portfolio.period_stats(history, [], None, "1M", 1).first == (
            _utc(2027, 3, 1, 5))
        later = _utc(2027, 3, 31, 12)
        times = (start, *live_portfolio.day_ends(start, later), later)
        history = History(times, (100.0,) * len(times), {}, {}, {}, {}, (0.0,) * len(times))
        assert live_portfolio.period_stats(history, [], None, "1M", 1).first == (
            _utc(2027, 3, 1, 5))

    def test_a_period_opens_at_a_close_not_at_a_deposit(self):
        # A month before Dec 1 10:00 New York time is Nov 1 10:00 (15:00 UTC); a
        # deposit landed after it, at 20:00 UTC, before the Nov 2 close
        times = (START, _utc(2026, 9, 29, 4), _utc(2026, 11, 1, 4), _utc(2026, 11, 1, 20),
                 _utc(2026, 11, 2, 5), _utc(2026, 12, 1, 15))
        history = History(times, (100.0, 100.0, 100.0, 105.0, 105.0, 105.0), {}, {}, {}, {},
                          (0.0, 0.0, 0.0, 5.0, 0.0, 0.0))
        assert live_portfolio.period_stats(history, [], None, "1M", 1).first == (
            _utc(2026, 11, 2, 5))
        assert live_portfolio.period_stats(history, [], None, "3M", 3).first == START

    def test_a_month_back_starts_at_the_first_close_after_it(self):
        read_at = _utc(2026, 12, 15, 12)
        times = _times(read_at)
        history = History(times, tuple(100.0 + i for i in range(len(times))), {}, {}, {}, {},
                          (0.0,) * len(times))
        returns = [TradeReturn("T1", T0, D(4), 0.1, D(4), True),
                   TradeReturn("T2", _utc(2026, 12, 1), D(2), -0.2, D(2), True)]
        stats = live_portfolio.period_stats(history, returns, None, "1M", 1)
        # A month before Dec 15 12:00 UTC is Nov 15 12:00; the next close is midnight
        # New York time on Nov 16 (05:00 UTC)
        assert stats.first == _utc(2026, 11, 16, 5)
        assert (stats.purchases, stats.pairs, stats.mean_trade) == (1, 2, -0.2)


class TestGroupReturns:
    def test_reinvested_money_is_not_put_in_twice(self):
        t = [T0 + timedelta(minutes=i) for i in range(4)]
        fills = [_fill("b1", "M1", t[0], True, 100, "0.10"),     # spend $10
                 _fill("s1", "M1", t[1], False, 100, "0.12"),    # get $12 back
                 _fill("b2", "M2", t[2], True, 100, "0.12"),     # spend the $12
                 _fill("s2", "M2", t[3], False, 100, "0.144")]   # get $14.40 back
        ledger = _ledger(fills, {"b1": "T1", "b2": "T2"},
                         legs={"T1": frozenset({"M1", "X1"}), "T2": frozenset({"M2", "X2"})})
        read_at = _utc(2026, 10, 1, 12)
        history = live_portfolio.build_history(_after(D(100), ledger, read_at=read_at), ledger,
                                               Marks({}, {}, read_at), _crypto, START)
        stats = live_portfolio.period_stats(history, [], None, "All", 0)
        [crypto] = stats.groups
        assert crypto.group == "Crypto"
        assert crypto.put_in == pytest.approx(10.0)
        assert crypto.pnl == pytest.approx(4.4)
        assert crypto.ret == pytest.approx(0.44)

    def test_value_held_at_the_period_start_is_put_in(self):
        times = _times(_utc(2026, 10, 1, 12))
        history = History(times, (0.0,) * 5, {"Crypto": (0.0, 20.0, 22.0, 25.0, 25.0)},
                          {"Crypto": (0.0, -20.0, -20.0, -20.0, -20.0)},
                          {"Crypto": (0.0, 20.0, 20.0, 20.0, 20.0)},
                          {"Crypto": ((T0, -20.0),)}, (0.0,) * 5)
        # From the first close on: $20 held then, nothing more put in, $5 gained
        stats = live_portfolio.period_stats(history, [], None, "1M", 1)
        [crypto] = live_portfolio._group_returns(history, 1)
        assert (crypto.put_in, crypto.pnl, crypto.ret) == (20.0, 5.0, 0.25)
        assert stats.groups[0].put_in == 20.0           # the whole history: $20 spent
        assert stats.groups[0].pnl == 5.0

    def test_nothing_put_in_has_no_return(self):
        times = _times(_utc(2026, 10, 1, 12))
        history = History(times, (0.0,) * 5, {OTHER_BETS: (0.0,) * 5},
                          {OTHER_BETS: (0.0, 1.0, 1.0, 1.0, 1.0)}, {OTHER_BETS: (0.0,) * 5},
                          {OTHER_BETS: ((T0, 1.0),)}, (0.0,) * 5)
        [group] = live_portfolio._group_returns(history, 0)
        assert (group.pnl, group.put_in, group.ret) == (1.0, 0.0, None)


class TestTradeReturns:
    TRADE = _bot("T1", [("A", "yes", 10), ("B", "no", 10)], T0 + timedelta(minutes=1))

    def _pair_ledger(self, *sales):
        """T1's pair (YES A at 0.30, NO B at 0.40), then your sales of YES A: (count, price)."""
        fills = [_fill("b", "B", T0, False, 10, "0.60"),
                 _fill("a", "A", T0 + timedelta(seconds=1), True, 10, "0.30")]
        fills += [_fill(f"sale{i}", "A", T0 + timedelta(days=1 + i), False, count, price)
                  for i, (count, price) in enumerate(sales)]
        return _ledger(fills, {"a": "T1", "b": "T1"}, legs={"T1": frozenset({"A", "B"})})

    def test_a_sale_counts_toward_its_own_trade(self):
        marks = Marks({}, {"A": D("0.55"), "B": D("0.55")}, READ)
        [ret] = live_portfolio.trade_returns([self.TRADE], self._pair_ledger((10, "0.50")),
                                             marks)
        # Spent $7; $5 back from the sale; 10 NO B worth $4.50 now
        assert ret.ret == pytest.approx((-7 + 5 + 4.5) / 7)
        assert (ret.trade_id, ret.opened, ret.pairs) == ("T1", T0, D(10))
        assert (ret.open_pairs, ret.still_open) == (D(0), True)

    def test_a_partly_sold_trade_counts_only_what_is_held(self):
        marks = Marks({}, {"A": D("0.55"), "B": D("0.55")}, READ)
        [ret] = live_portfolio.trade_returns([self.TRADE], self._pair_ledger((4, "0.50")),
                                             marks)
        assert ret.open_pairs == D(6) and ret.still_open
        assert ret.ret == pytest.approx((-7 + 2 + 6 * 0.55 + 4.5) / 7)

    def test_contracts_never_priced_count_at_cost(self):
        [ret] = live_portfolio.trade_returns([self.TRADE], self._pair_ledger(),
                                             Marks({}, {}, READ))
        assert ret.ret == 0.0 and ret.open_pairs == D(10)

    def test_a_fully_closed_trade_is_no_longer_open(self):
        fills = [_fill("b", "B", T0, False, 10, "0.60"),
                 _fill("a", "A", T0 + timedelta(seconds=1), True, 10, "0.30")]
        payouts = [Payout("A", T0 + timedelta(days=2), "yes", None, D(10)),
                   Payout("B", T0 + timedelta(days=2), "yes", None, D(0))]
        ledger = _ledger(fills, {"a": "T1", "b": "T1"}, payouts,
                         legs={"T1": frozenset({"A", "B"})})
        [ret] = live_portfolio.trade_returns([self.TRADE], ledger, Marks({}, {}, READ))
        assert ret.ret == pytest.approx(3 / 7)
        assert (ret.open_pairs, ret.still_open) == (D(0), False)

    def test_rolled_back_rows_and_one_leg_purchases_are_left_out(self):
        logged = T0 + timedelta(minutes=1)
        trades = [self.TRADE,
                  _bot("T2", [("D", "yes", 5), ("C", "no", 5)], logged, status="rolled_back"),
                  _bot("T3", [("E", "yes", 5), ("F", "no", 5)], logged)]
        fills = [_fill("b", "B", T0, False, 10, "0.60"),
                 _fill("a", "A", T0 + timedelta(seconds=1), True, 10, "0.30"),
                 _fill("r-no", "C", T0, False, 5, "0.50"),
                 _fill("r-back", "C", T0 + timedelta(seconds=2), True, 5, "0.55"),
                 _fill("t3-a", "E", T0, True, 5, "0.40")]
        owners = {"a": "T1", "b": "T1", "r-no": "T2", "r-back": "T2", "t3-a": "T3"}
        legs = {t.trade_id: frozenset(leg.ticker for leg in t.legs) for t in trades}
        returns = live_portfolio.trade_returns(trades, _ledger(fills, owners, legs=legs),
                                               Marks({}, {}, READ))
        assert [r.trade_id for r in returns] == ["T1"]

    @pytest.mark.parametrize("status, counted", [
        ("executed", True), ("manual_review", True),
        ("rolled_back", False), ("rollback_failed", False)])
    def test_only_a_row_that_can_hold_a_pair_counts_when_both_orders_are_found(self, status,
                                                                                counted):
        # Both orders of the row are its own in the ledger; a manual-review row
        # whose two orders filled holds a pair, a rolled-back or failed-unwind
        # row never does (its YES leg never filled), whatever the ledger says
        trade = _bot("T1", [("A", "yes", 10), ("B", "no", 10)], T0 + timedelta(minutes=1),
                     status=status)
        ledger = self._pair_ledger()
        returns = live_portfolio.trade_returns([trade], ledger, Marks({}, {}, READ))
        assert [(r.trade_id, r.pairs, r.open_pairs) for r in returns] == (
            [("T1", D(10), D(10))] if counted else [])


class TestPairWeighted:
    @pytest.mark.parametrize("seed", range(200))
    def test_it_matches_numpy_on_the_list_with_each_pair_once(self, seed):
        rng = random.Random(seed)
        returns = [TradeReturn(f"T{i}", T0, D(rng.randint(1, 6)),
                               rng.choice([rng.uniform(-1, 2), round(rng.uniform(-1, 1), 1)]),
                               D(0), False)
                   for i in range(rng.randint(1, 12))]
        expanded = [r.ret for r in returns for _ in range(int(r.pairs))]
        mean, median = live_portfolio.pair_weighted(returns)
        assert median == np.median(expanded)
        assert mean == pytest.approx(float(np.mean(expanded)), rel=1e-12)

    def test_no_purchase_has_no_figure(self):
        assert live_portfolio.pair_weighted([]) == (None, None)


# ---- the cash each run logged ----------------------------------------------

class TestCheckLoggedCash:
    def _setup(self, shards=1):
        """$3.00 of the bot's YES A bought at T0, leaving $20.2275 of cash now."""
        fill = _fill("bot", "A", T0, True, 10, "0.30")
        ledger = _ledger([fill], {"bot": "T1"})
        account = Account(READ, D("20.2275"), shards, None, {"A": D(10)}, (fill,), (), (), False)
        return account, ledger

    def _check(self, banner, *, shards=1, first=T0, logged=T0 + timedelta(minutes=1)):
        account, ledger = self._setup(shards)
        run = RunStart(logged - timedelta(hours=6), logged, D(banner))
        fills = {run.logged_at: first} if first else {}
        return live_portfolio.check_logged_cash(account, ledger, [run], fills)

    def test_a_banner_rounded_down_to_the_cent_matches(self):
        # Rebuilt just before the run's first fill: $23.2275, logged as $23.22
        check = self._check("23.22")
        assert (check.matched, check.checked, check.worst, check.misses) == (1, 1, D("0.0075"), ())

    @pytest.mark.parametrize("banner, shards, matches", [
        ("23.2075", 1, False),     # $0.02 more rebuilt than logged
        ("23.2475", 1, False),     # $0.02 less
        ("23.2325", 1, True),      # half a cent less: still a match
        ("23.2326", 1, False),
        ("23.2175", 1, False),     # a whole cent more, on one shard
        ("23.2175", 2, True),      # ... is within two shards' rounding
        ("23.2075", 2, False),
    ])
    def test_the_bounds(self, banner, shards, matches):
        check = self._check(banner, shards=shards)
        assert check.matched == int(matches)
        logged = T0 + timedelta(minutes=1)
        assert check.misses == (() if matches else ((logged, D(banner), D("23.2275")),))

    def test_a_run_with_no_fills_is_checked_at_its_log_time(self):
        # After the fill: $20.2275 rebuilt at the log time
        assert self._check("20.22", first=None).matched == 1
        assert self._check("23.22", first=None).matched == 0

    def test_a_run_before_the_bots_first_fill_is_checked_too(self):
        # A real run whose orders never filled, logged an hour before the
        # bot's first fill: its cash is rebuilt at its log time
        before = T0 - timedelta(hours=1)
        check = self._check("23.22", first=None, logged=before)
        assert (check.checked, check.matched) == (1, 1)
        miss = self._check("23.20", first=None, logged=before)
        assert miss.misses == ((before, D("23.20"), D("23.2275")),)

    def test_a_miss_is_one_sentence_with_the_rebuilt_cash_to_the_hundredth_of_a_cent(self):
        logged = _utc(2026, 9, 28, 7, 11, 35)
        check = live_portfolio.CashCheck(0, 2, D("1.004999999999"), (
            (logged, D("101.00"), D("100.005000000000")),
            (logged + timedelta(days=1), D("50.10"), D("49.123456789012"))))
        assert live_portfolio._cash_miss_warnings(check) == [
            "The run logged at 2026-09-28 07:11 UTC wrote $101.00 as its cash before trading; "
            "Kalshi's records rebuild $100.0050 then",
            "The run logged at 2026-09-29 07:11 UTC wrote $50.10 as its cash before trading; "
            "Kalshi's records rebuild $49.1235 then"]


# ---- the whole view, and its log line ----------------------------------------

def _view(**changes) -> LiveView:
    """A LiveView built from Decimal fixtures, as build_live_view returns one."""
    holding = Holding("A", "Will A?", "Crypto", "yes", D("10"), D("0.45"), D("4.50"), D("3.0700"))
    view = LiveView(READ, D("20.2275"), D("7.85"), (holding,), ("Crypto",), None, None, (),
                    live_portfolio.CashCheck(1, 1, D("0.0075"), ()), None, False, ())
    return dataclasses.replace(view, **changes)


class TestSnapshot:
    def test_the_record_is_plain_json(self):
        record = live_portfolio.snapshot_record(_view())
        assert json.loads(json.dumps(record)) == {
            "read_at": "2026-10-07T12:00:00+00:00", "cash": "20.2275",
            "kalshi_positions_value": "7.85", "holdings_value": "4.50", "changing": False,
            "holdings": [{"ticker": "A", "title": "Will A?", "group": "Crypto", "side": "yes",
                          "contracts": "10", "price": "0.45", "value": "4.50",
                          "cost": "3.0700", "subtitle": "", "event_ticker": ""}]}
        never_priced = dataclasses.replace(_view().holdings[0], price=None)
        record = live_portfolio.snapshot_record(
            _view(holdings=(never_priced,), kalshi_positions_value=None))
        assert json.loads(json.dumps(record))["holdings"][0]["price"] is None
        assert record["kalshi_positions_value"] is None

    def test_each_read_adds_one_line(self, tmp_path):
        log = config.LIVE_PORTFOLIO_LOG_FILE
        assert log.parent == tmp_path
        live_portfolio.append_snapshot(_view())
        live_portfolio.append_snapshot(_view(cash=D("1")))
        lines = log.read_text(encoding="utf-8").splitlines()
        assert [json.loads(line)["cash"] for line in lines] == ["20.2275", "1"]

    def test_a_write_that_fails_is_a_warning(self, tmp_path, monkeypatch, caplog):
        monkeypatch.setattr(config, "LIVE_PORTFOLIO_LOG_FILE", tmp_path)   # a folder
        with caplog.at_level(logging.WARNING):
            live_portfolio.append_snapshot(_view())
        assert [r.levelno for r in caplog.records] == [logging.WARNING]
        assert "Could not add this read of the account to" in caplog.text

    def test_a_record_json_cannot_hold_is_a_warning(self, caplog):
        odd = dataclasses.replace(_view().holdings[0], group=object())
        with caplog.at_level(logging.WARNING):
            live_portfolio.append_snapshot(_view(holdings=(odd,)))
        assert "TypeError" in caplog.text
        assert not config.LIVE_PORTFOLIO_LOG_FILE.exists()


def _market_row(ticker, event, bid, ask) -> dict:
    """A market record as /markets sends it."""
    return {"ticker": ticker, "event_ticker": event, "title": f"Will {ticker}?",
            "status": "active", "result": "", "yes_bid_dollars": bid, "yes_ask_dollars": ask,
            "last_price_dollars": bid}


@pytest.fixture
def whole_account(kalshi, monkeypatch):
    """FakeKalshi for the account and its markets, FakeCandles for the daily candles."""
    fake = FakeCandles()

    def get(client, path, **params):
        if path.endswith("candlesticks"):
            return fake(client, path, **params)
        return kalshi(client, path, **params)

    monkeypatch.setattr(historical, "_historical_get", get)
    return kalshi, fake


class TestBuildLiveView:
    CATEGORIES = {"KXCRYPTO": ("Crypto", ("BTC",))}

    def _account(self, kalshi, fake, positions=None):
        """One bot pair (13 YES A1 at 0.30, 13 NO B1 at 0.40) and 5 YES X of your own."""
        kalshi.rows["/portfolio/fills"] = [
            fill_row("bot-a", "A1", T0 - timedelta(seconds=2), "bid", 13, "0.30"),
            fill_row("bot-b", "B1", T0 - timedelta(seconds=3), "ask", 13, "0.60"),
            fill_row("mine", "X", T0 - timedelta(days=2), "bid", 5, "0.20"),
        ]
        held = positions or {"A1": "13.00", "B1": "-13.00", "X": "5.00"}
        kalshi.rows["/portfolio/positions"] = [{"ticker": t, "position_fp": n}
                                               for t, n in held.items()]
        # $100.0050 before the bot's $9.10 pair
        kalshi.balance = {"balance_breakdown": [{"exchange_index": 0,
                                                 "balance_dollars": "90.9050"}],
                          "portfolio_value": 1225}
        kalshi.rows["/markets"] = [{**_market_row("A1", "KXCRYPTO-26DEC", "0.40", "0.50"),
                                    "yes_sub_title": "Above $100k"},
                                   _market_row("B1", "KXCRYPTO-26DEC31", "0.55", "0.65"),
                                   _market_row("X", "KXWEATHER-26OCT", "0.20", "0.30")]
        closes = live_portfolio.day_ends(T0 - timedelta(days=3), datetime.now(UTC))
        for ticker, mid in (("A1", ("0.40", "0.50")), ("B1", ("0.55", "0.65")),
                            ("X", ("0.20", "0.30"))):
            fake.batch[ticker] = [_candle(c, *mid) for c in closes]

    def test_a_bot_pair_and_your_own_bet(self, whole_account, log_paths, pacific, monkeypatch):
        kalshi, fake = whole_account
        self._account(kalshi, fake)
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed")], 100.00)
        view = live_portfolio.build_live_view(
            object(), risk_free=None, series_categories=self.CATEGORIES,
            trade_logs=live_portfolio.trade_log_paths())
        assert view.warnings == ()
        assert (view.cash, view.kalshi_positions_value) == (D("90.9050"), D("12.25"))
        assert view.group_order == ("Crypto", OTHER_BETS)
        assert [(h.ticker, h.group, h.side, h.contracts, h.price, h.value, h.cost)
                for h in view.holdings] == [
            ("A1", "Crypto", "yes", D(13), D("0.45"), D("5.85"), D("3.9000")),
            ("B1", "Crypto", "no", D(13), D("0.40"), D("5.20"), D("5.2000")),
            ("X", OTHER_BETS, "yes", D(5), D("0.25"), D("1.25"), D("1.0000"))]
        assert view.holdings[0].title == "Will A1?"
        assert [(h.subtitle, h.event_ticker) for h in view.holdings] == [
            ("Above $100k", "KXCRYPTO-26DEC"), ("", "KXCRYPTO-26DEC31"), ("", "KXWEATHER-26OCT")]
        assert view.holdings_value == D("12.30")
        # The banner's $100.00 against $100.0050 rebuilt just before the pair
        assert (view.cash_check.matched, view.cash_check.checked) == (1, 1)
        assert view.history.times[0] == T0 - timedelta(seconds=3, microseconds=1)
        assert [p.label for p in view.periods] == [label for label, _ in
                                                   config.LIVE_DASHBOARD_PERIODS]
        [trade] = view.trades
        assert trade.ret == pytest.approx((-9.10 + 5.85 + 5.20) / 9.10)
        assert view.periods[0].purchases == 1
        # The history's last moment is the read, valued as the holdings are
        assert view.history.total(-1) == pytest.approx(90.905 + 12.30)
        assert {t for p in fake.batches() for t in p["market_tickers"].split(",")} == {
            "A1", "B1", "X"}
        assert view.changing is False

    def _two_categories(self, kalshi, fake):
        """
        Two bot pairs in one run, and a bet of your own held at the start.

        Crypto: 13 YES A1 at 0.30 and 13 NO B1 at 0.40, $9.10. Climate: 10 YES
        C1 at 0.20 (a weather market) and 10 NO D1 at 0.20 (a crypto market:
        a purchase's category is its market A's), $4.00, but worth more than
        Crypto now. Your 5 YES X, bought two days before the bot's first
        fill, paid out a day after it; Kalshi lists that settlement and no
        position, so only the ledger shows X held at the start.
        """
        self._account(kalshi, fake, positions={"A1": "13.00", "B1": "-13.00", "C1": "10.00",
                                               "D1": "-10.00"})
        kalshi.rows["/portfolio/fills"] += [
            fill_row("bot-c", "C1", T0 - timedelta(seconds=5), "bid", 10, "0.20"),
            fill_row("bot-d", "D1", T0 - timedelta(seconds=6), "ask", 10, "0.80")]
        paid = T0 + timedelta(days=1)
        kalshi.rows["/portfolio/settlements"] = [
            {"ticker": "X", "settled_time": _iso(paid), "market_result": "yes", "revenue": 500}]
        # $100.0050 before the bot's two pairs ($13.10), and X's $5 since
        kalshi.balance = {"balance_breakdown": [{"exchange_index": 0,
                                                 "balance_dollars": "91.9050"}],
                          "portfolio_value": 2905}
        kalshi.rows["/markets"] = [
            *(row for row in kalshi.rows["/markets"] if row["ticker"] != "X"),
            _market_row("C1", "KXWEATHER-26OCT", "0.85", "0.95"),
            _market_row("D1", "KXCRYPTO-26NOV", "0.05", "0.15"),
            {**_market_row("X", "KXWEATHER-26OCT", "0.98", "1.00"), "status": "finalized",
             "result": "yes", "settlement_ts": _iso(paid)}]
        closes = live_portfolio.day_ends(T0 - timedelta(days=3), datetime.now(UTC))
        fake.batch["C1"] = [_candle(c, "0.85", "0.95") for c in closes]
        fake.batch["D1"] = [_candle(c, "0.05", "0.15") for c in closes]
        fake.batch["X"] = [_candle(c, "0.20", "0.30") for c in closes if c <= paid]

    def test_categories_are_market_as_and_ordered_by_first_purchase(self, whole_account,
                                                                     log_paths, pacific,
                                                                     monkeypatch):
        kalshi, fake = whole_account
        self._two_categories(kalshi, fake)
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed"),
                                     _trade_result("C1", "D1", "executed", count=10)], 100.00)
        view = live_portfolio.build_live_view(
            object(), risk_free=None,
            series_categories={**self.CATEGORIES, "KXWEATHER": ("Climate", ("Rain",))},
            trade_logs=live_portfolio.trade_log_paths())
        assert view.warnings == ()
        # Climate's first fill (D1, 6 s before the run) came before Crypto's (B1,
        # 3 s before), so Climate comes first though Crypto put in more ($9.10
        # against $4.00); within a group, by event (D1's KXCRYPTO-26NOV before
        # C1's KXWEATHER-26OCT)
        assert view.group_order == ("Climate", "Crypto", OTHER_BETS)
        assert [(h.ticker, h.group, h.value) for h in view.holdings] == [
            ("D1", "Climate", D("9.00")), ("C1", "Climate", D("9.00")),
            ("A1", "Crypto", D("5.85")), ("B1", "Crypto", D("5.20"))]
        # X was looked up once the ledger showed it held at the start, so it
        # has its daily price there (0.25), not what it cost (0.20)
        asked = [params["tickers"] for path, params in kalshi.calls if path == "/markets"]
        assert asked == ["A1,B1,C1,D1", "X"]
        assert view.history.value[OTHER_BETS][0] == pytest.approx(1.25)
        assert (view.cash_check.matched, view.cash_check.checked) == (1, 1)

    def test_an_account_that_kept_changing_says_so(self, whole_account, log_paths, pacific,
                                                    monkeypatch):
        kalshi, fake = whole_account
        self._account(kalshi, fake)
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed")], 100.00)
        read = live_portfolio.read_account
        monkeypatch.setattr(live_portfolio, "read_account",
                            lambda client: dataclasses.replace(read(client), changing=True))
        view = live_portfolio.build_live_view(
            object(), risk_free=None, series_categories=self.CATEGORIES,
            trade_logs=live_portfolio.trade_log_paths())
        assert view.changing is True
        assert view.warnings == ("The account kept changing while it was read: the figures "
                                 "may not all be from one moment",)

    def test_every_real_run_is_checked_and_a_miss_is_a_warning(self, whole_account, log_paths,
                                                              pacific, monkeypatch):
        kalshi, fake = whole_account
        self._account(kalshi, fake)
        # A run two hours before the bot's first fill, whose pair failed and sent
        # nothing that filled, then the bot's real pair, logged with a wrong cash
        _write_run(monkeypatch, T0 - timedelta(hours=2), [_trade_result("A1", "B1", "failed")],
                   100.00)
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed")], 101.00)
        view = live_portfolio.build_live_view(
            object(), risk_free=None, series_categories=self.CATEGORIES,
            trade_logs=live_portfolio.trade_log_paths())
        assert (view.cash_check.matched, view.cash_check.checked) == (1, 2)
        assert view.warnings == ("The run logged at 2026-09-28 09:00 UTC wrote $101.00 as its "
                                 "cash before trading; Kalshi's records rebuild $100.0050 then",)

    def test_kalshi_holding_something_else_is_a_warning(self, whole_account, log_paths, pacific,
                                                         monkeypatch):
        kalshi, fake = whole_account
        self._account(kalshi, fake, positions={"A1": "12.00", "B1": "-13.00", "X": "5.00",
                                               "Y": "-2.00"})
        _write_run(monkeypatch, T0, [_trade_result("A1", "B1", "executed")], 100.00)
        view = live_portfolio.build_live_view(
            object(), risk_free=None, series_categories=self.CATEGORIES,
            trade_logs=live_portfolio.trade_log_paths())
        assert view.warnings == (
            "A1: Kalshi holds 12 YES, but the fills and payouts add up to 13 YES",
            "Y: Kalshi holds 2 NO, but the fills and payouts add up to none")

    def test_before_the_bots_first_trade_there_is_no_history(self, whole_account, log_paths):
        kalshi, fake = whole_account
        self._account(kalshi, fake)
        view = live_portfolio.build_live_view(object(), risk_free=None, series_categories=None,
                                              trade_logs=live_portfolio.trade_log_paths())
        assert (view.history, view.periods, view.trades) == (None, None, ())
        assert view.group_order == (OTHER_BETS,)
        assert [(h.ticker, h.value) for h in view.holdings] == [
            ("A1", D("5.85")), ("B1", D("5.20")), ("X", D("1.25"))]
        assert fake.calls == []                          # no daily prices to read
        assert view.cash_check.checked == 0

    def test_a_later_bigger_category_does_not_take_an_earlier_ones_place(self):
        # Each category's first change is its first purchase: Early before
        # Middle before Late, though Late put in and is worth the most, and
        # Other bets last, though your own bet changed before Middle and Late
        times = _times(_utc(2026, 10, 1, 12))
        first = {"Early": T0, "Middle": T0 + timedelta(hours=1),
                 "Late": T0 + timedelta(days=2), OTHER_BETS: T0 + timedelta(minutes=1)}
        spent = {"Early": 2.0, "Middle": 3.0, "Late": 500.0, OTHER_BETS: 1.0}
        history = History(
            times, (0.0,) * 5, {g: (0.0, 0.0, s, s, s) for g, s in spent.items()},
            {g: (0.0, 0.0, -s, -s, -s) for g, s in spent.items()},
            {g: (0.0, 0.0, s, s, s) for g, s in spent.items()},
            {g: ((when, -spent[g]),) for g, when in first.items()}, (0.0,) * 5)
        before = live_portfolio._group_order(dataclasses.replace(
            history, **{name: {g: v for g, v in getattr(history, name).items() if g != "Late"}
                        for name in ("value", "net_cash", "spent", "steps")}), [])
        assert before == ("Early", "Middle", OTHER_BETS)
        assert live_portfolio._group_order(history, []) == (
            "Early", "Middle", "Late", OTHER_BETS)
        # Two that first appeared at one moment go by name; a group held now
        # with no change in the history comes after those that have one
        tied = dataclasses.replace(history, steps={**history.steps,
                                                   "Late": ((T0, -500.0),)})
        newcomer = Holding("N", "Will N?", "Aardvark", "yes", D(1), D("0.5"), D("0.5"), D(1))
        assert live_portfolio._group_order(tied, [newcomer]) == (
            "Early", "Late", "Middle", "Aardvark", OTHER_BETS)

    def test_the_rungs_of_one_question_sit_together(self):
        def held(ticker, event, group, value, side="yes"):
            return Holding(ticker, f"Will {ticker}?", group, side, D(1), D("0.5"), D(value),
                           D(1), "", event)

        rung_dec = held("Q-26DEC01", "Q", "Tech", "1", "no")
        rung_nov = held("Q-26NOV01", "Q", "Tech", "9")
        other = held("P-1", "P", "Tech", "5")
        mine = held("Z", "Z", OTHER_BETS, "50")
        ordered = live_portfolio._sorted_holdings([mine, rung_nov, other, rung_dec],
                                                  ("Tech", OTHER_BETS))
        # By value the two rungs would sit apart, with P-1 ($5) between them
        assert [h.ticker for h in ordered] == ["P-1", "Q-26DEC01", "Q-26NOV01", "Z"]

    def test_a_purchase_whose_market_was_not_found_is_filed_by_its_ticker(self):
        markets = {"A": _market("A", event="KXFED-26DEC")}
        assert live_portfolio._event_of("A", markets) == "KXFED-26DEC"
        assert live_portfolio._event_of("KXBTC-26OCT-T50", {}) == "KXBTC-26OCT"
        assert live_portfolio._event_of("PLAIN", {}) == "PLAIN"
        unlisted = dataclasses.replace(_market("B-1"), event_ticker="")
        assert live_portfolio._event_of("B-1", {"B-1": unlisted}) == "B"
